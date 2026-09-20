from __future__ import annotations

import base64
import copy
import hashlib
import json
import os
import re
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path

from PySide6.QtGui import (
    QColor,
    QFont,
    QImage,
    QPen,
    QTextBlockFormat,
    QTextCharFormat,
    QTextCursor,
    QTextDocument,
    QTextFormat,
    QTextFrameFormat,
    QTextImageFormat,
    QTextListFormat,
    QTextLength,
    QTextTable,
    QTextTableFormat,
)
from PySide6.QtCore import QRectF, Qt

import docx
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt, RGBColor

from .editor import (COLUMN_BREAK_MARKER, COLUMN_BREAK_URL, SafeDocument,
                     STYLE_PRESETS, block_break_kind, mm_to_px)
from .models import DocumentState, PageSettings

MAX_ZIP_ENTRIES = 4096
MAX_UNCOMPRESSED_BYTES = 128 * 1024 * 1024
MAX_METADATA_BYTES = 16 * 1024 * 1024
EMU_PER_PX = 9525
PT_PER_PX = 72.0 / 96.0
META_PART = "customXml/folio-state.json"
META_CONTENT_TYPE = "application/vnd.folio.document+json"
META_REL_TYPE = "urn:folio:document-state"

IMPORT_WARNING = (
    "This DOCX was imported with best effort: complex sections, fields, "
    "floating graphics, or Word-native tracked changes may not be retained."
)


class DocxError(RuntimeError):
    pass


def _word_digest(members: dict[str, bytes]) -> str:
    digest = hashlib.sha256()
    for name in sorted(n for n in members if n.startswith("word/")):
        digest.update(name.encode("utf-8"))
        digest.update(b"\0")
        digest.update(members[name])
    return digest.hexdigest()


def atomic_write(path: Path, data: bytes) -> None:
    path = Path(path)
    fd, temp_name = tempfile.mkstemp(prefix=path.name + ".", suffix=".tmp",
                                     dir=str(path.parent))
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(data)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temp_name, path)
    except BaseException:
        try:
            os.unlink(temp_name)
        except OSError:
            pass
        raise


def _px_to_pt(px: float) -> float:
    return px * PT_PER_PX


_ALIGNMENT_MAP = {
    Qt.AlignmentFlag.AlignLeft: WD_ALIGN_PARAGRAPH.LEFT,
    Qt.AlignmentFlag.AlignRight: WD_ALIGN_PARAGRAPH.RIGHT,
    Qt.AlignmentFlag.AlignHCenter: WD_ALIGN_PARAGRAPH.CENTER,
    Qt.AlignmentFlag.AlignJustify: WD_ALIGN_PARAGRAPH.JUSTIFY,
}

_HIGHLIGHT_COLORS = [
    ("#ffff00", WD_COLOR_INDEX.YELLOW),
    ("#00ff00", WD_COLOR_INDEX.BRIGHT_GREEN),
    ("#00ffff", WD_COLOR_INDEX.TURQUOISE),
    ("#ff00ff", WD_COLOR_INDEX.PINK),
    ("#ff0000", WD_COLOR_INDEX.RED),
    ("#0000ff", WD_COLOR_INDEX.BLUE),
    ("#008080", WD_COLOR_INDEX.TEAL),
    ("#008000", WD_COLOR_INDEX.GREEN),
    ("#800080", WD_COLOR_INDEX.VIOLET),
    ("#808080", WD_COLOR_INDEX.GRAY_50),
    ("#000000", WD_COLOR_INDEX.BLACK),
    ("#ffffff", WD_COLOR_INDEX.WHITE),
]


def _nearest_highlight(color: QColor):
    best = None
    best_dist = float("inf")
    r, g, b = color.red(), color.green(), color.blue()
    for hex_value, index in _HIGHLIGHT_COLORS:
        c = QColor(hex_value)
        dist = (r - c.red()) ** 2 + (g - c.green()) ** 2 + (b - c.blue()) ** 2
        if dist < best_dist:
            best_dist = dist
            best = index
    return best


def _preset_for_block(block) -> str | None:
    from .design import STYLE_PROPERTY, style_name

    if block.blockFormat().hasProperty(STYLE_PROPERTY):
        return style_name(block)
    level = block.blockFormat().headingLevel()
    if level > 0:
        return f"Heading {min(level, 3)}"
    fmt = block.charFormat()
    size = fmt.fontPointSize()
    color = fmt.foreground().color().name() if fmt.hasProperty(QTextFormat.Property.ForegroundBrush) else ""
    for name in ("Title", "Subtitle", "Quote"):
        preset = STYLE_PRESETS[name]
        if abs(size - preset.size_pt) < 0.6:
            if name == "Quote":
                if block.blockFormat().leftMargin() > 1:
                    return name
                continue
            return name
    return None


def _decode_image_name(name: str) -> QImage | None:
    if not name.startswith("data:image/"):
        return None
    marker = ";base64,"
    idx = name.find(marker)
    if idx == -1:
        return None
    try:
        payload = base64.b64decode(name[idx + len(marker):])
        image = QImage.fromData(payload)
        return image if not image.isNull() else None
    except Exception:
        return None


class _NumberingRegistry:
    def __init__(self):
        self._by_list: dict[int, int] = {}
        self._abstract: list[tuple[int, str]] = []
        self._num_to_abstract: dict[int, int] = {}
        self._next_num = 1
        self._next_abstract = 0

    def num_id(self, text_list) -> int:
        from shiboken6 import getCppPointer

        key = getCppPointer(text_list)[0]
        existing = self._by_list.get(key)
        if existing is not None:
            return existing
        fmt = text_list.format()
        number_format = {
            QTextListFormat.Style.ListDecimal: "decimal",
            QTextListFormat.Style.ListLowerAlpha: "lowerLetter",
            QTextListFormat.Style.ListUpperAlpha: "upperLetter",
            QTextListFormat.Style.ListLowerRoman: "lowerRoman",
            QTextListFormat.Style.ListUpperRoman: "upperRoman",
        }.get(fmt.style(), "bullet")
        num_id = self._next_num
        self._next_num += 1
        self._by_list[key] = num_id
        self._abstract.append((self._next_abstract, number_format))
        self._num_to_abstract[num_id] = self._next_abstract
        self._next_abstract += 1
        return num_id

    def numbering_xml(self) -> bytes | None:
        if not self._abstract:
            return None
        w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        parts = [f'<w:numbering {w}>']
        for abstract_id, number_format in self._abstract:
            parts.append(f'<w:abstractNum w:abstractNumId="{abstract_id}">')
            for level in range(9):
                ordered = number_format != "bullet"
                num_fmt = number_format
                lvl_text = f"%{level + 1}." if ordered else ""
                font = "" if ordered else '<w:rPr><w:rFonts w:ascii="Symbol" w:hAnsi="Symbol" w:hint="default"/></w:rPr>'
                left = 720 * (level + 1)
                parts.append(
                    f'<w:lvl w:ilvl="{level}"><w:start w:val="1"/>'
                    f'<w:numFmt w:val="{num_fmt}"/><w:lvlText w:val="{lvl_text}"/>'
                    f'<w:lvlJc w:val="left"/>'
                    f'<w:pPr><w:ind w:left="{left}" w:hanging="360"/></w:pPr>'
                    f'{font}</w:lvl>'
                )
            parts.append('</w:abstractNum>')
        for num_id, abstract_id in sorted(self._num_to_abstract.items()):
            parts.append(
                f'<w:num w:numId="{num_id}">'
                f'<w:abstractNumId w:val="{abstract_id}"/></w:num>'
            )
        parts.append('</w:numbering>')
        return "".join(parts).encode("utf-8")


class _Exporter:
    def __init__(self, state: DocumentState, document: QTextDocument):
        self.state = state
        self.src = document
        self.numbering = _NumberingRegistry()
        self.docx = docx.Document()
        from .document_features import node_ranges, section_ranges
        from .design import DesignSettings

        self.design = DesignSettings.from_dict(state.design)
        self._sections = section_ranges(document, state)
        self._ranges = node_ranges(document)
        self._bookmarks = {key: (index, "Folio_" + str(index))
                           for index, key in enumerate(self._ranges)}
        self._emitted_fields = set()
        self._footnotes = {}
        self._source_paragraphs = []
        normal = self.docx.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.font.color.rgb = RGBColor(0x26, 0x32, 0x44)
        self._export_styles()

    def build(self) -> bytes:
        self._export_sections()
        self._export_footnotes()
        self._export_body()
        self._export_floating_objects()
        buffer = BytesIO()
        self.docx.save(buffer)
        return buffer.getvalue()

    def _export_footnotes(self) -> None:
        from lxml import etree
        from .references import note_entries

        entries = note_entries(self.src, self.state, "footnote")
        if not entries:
            return
        root = OxmlElement("w:footnotes")
        for identifier, kind in ((-1, "separator"), (0, "continuationSeparator")):
            note = OxmlElement("w:footnote")
            note.set(qn("w:id"), str(identifier))
            note.set(qn("w:type"), kind)
            paragraph = OxmlElement("w:p")
            run = OxmlElement("w:r")
            run.append(OxmlElement("w:" + kind))
            paragraph.append(run)
            note.append(paragraph)
            root.append(note)
        for _position, key, number, node in entries:
            self._footnotes[key] = number
            note = OxmlElement("w:footnote")
            note.set(qn("w:id"), str(number))
            for index, line in enumerate(node["text"].split("\n")):
                paragraph = OxmlElement("w:p")
                if index == 0:
                    marker = OxmlElement("w:r")
                    marker.append(OxmlElement("w:footnoteRef"))
                    paragraph.append(marker)
                run = OxmlElement("w:r")
                text = OxmlElement("w:t")
                text.set(qn("xml:space"), "preserve")
                text.text = (" " if index == 0 else "") + line
                run.append(text)
                paragraph.append(run)
                note.append(paragraph)
            root.append(note)
        payload = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)
        part = Part(PackURI("/word/footnotes.xml"),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
                    payload, self.docx.part.package)
        self.docx.part.relate_to(part, RT.FOOTNOTES)

    def _export_styles(self) -> None:
        from .design import STYLE_SETS, resolved_style, theme_tokens
        from lxml import etree

        for name in STYLE_SETS[self.design.style_set]:
            style = self.docx.styles[name]
            values = resolved_style(name, self.design)
            style.font.name = values["font"]
            style.font.size = Pt(values["size"])
            style.font.bold = values["weight"] >= QFont.Weight.Bold
            style.font.italic = values["italic"]
            style.font.color.rgb = RGBColor.from_string(values["color"][1:])
            style.paragraph_format.space_before = Pt(_px_to_pt(values["before"]))
            style.paragraph_format.space_after = Pt(_px_to_pt(values["after"]))
        tokens = theme_tokens(self.design)
        theme = self.docx.part.part_related_by(RT.THEME)
        root = etree.fromstring(theme.blob)
        for tag, family in (("majorFont", tokens["heading_font"]),
                            ("minorFont", tokens["body_font"])):
            latin = root.find(".//" + qn("a:" + tag) + "/" + qn("a:latin"))
            if latin is not None:
                latin.set("typeface", family)
        for index, color in enumerate(tokens["accents"], start=1):
            accent = root.find(".//" + qn("a:accent" + str(index)))
            if accent is not None:
                for child in list(accent):
                    accent.remove(child)
                rgb = OxmlElement("a:srgbClr")
                rgb.set("val", color[1:])
                accent.append(rgb)
        theme._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        if self.design.page_color != "#ffffff":
            background = OxmlElement("w:background")
            background.set(qn("w:color"), self.design.page_color[1:])
            self.docx._element.insert(0, background)

    def _export_sections(self, section=None, settings=None) -> None:
        settings = settings or self._sections[0][2]
        width_mm, height_mm = settings.size_mm()
        section = section or self.docx.sections[0]
        section.page_width = Mm(width_mm)
        section.page_height = Mm(height_mm)
        section.orientation = (
            WD_ORIENT.LANDSCAPE if settings.landscape else WD_ORIENT.PORTRAIT
        )
        section.top_margin = Mm(settings.top_mm)
        section.right_margin = Mm(settings.right_mm)
        section.bottom_margin = Mm(settings.bottom_mm)
        section.left_margin = Mm(settings.left_mm)
        sect_pr = section._sectPr
        cols = sect_pr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            sect_pr.append(cols)
        cols.set(qn("w:num"), str(settings.columns))
        cols.set(qn("w:space"), str(int(settings.gutter_mm * 56.6929)))
        cols.set(qn("w:equalWidth"), "1")
        if self.design.page_border != "none":
            borders = OxmlElement("w:pgBorders")
            borders.set(qn("w:offsetFrom"), "page")
            style = {"box": "single", "double": "double", "dashed": "dashed"}[
                self.design.page_border]
            for side in ("top", "left", "bottom", "right"):
                border = OxmlElement("w:" + side)
                border.set(qn("w:val"), style)
                border.set(qn("w:sz"), str(round(self.design.page_border_width_pt * 8)))
                border.set(qn("w:space"), str(min(31, round(self.design.page_border_inset_mm * 72 / 25.4))))
                border.set(qn("w:color"), self.design.page_border_color[1:])
                borders.append(border)
            previous = sect_pr.find(qn("w:pgBorders"))
            if previous is not None:
                sect_pr.remove(previous)
            sect_pr.append(borders)
        line_numbers = sect_pr.find(qn("w:lnNumType"))
        if settings.line_numbering == "none":
            if line_numbers is not None:
                sect_pr.remove(line_numbers)
        else:
            if line_numbers is None:
                line_numbers = OxmlElement("w:lnNumType")
                sect_pr.append(line_numbers)
            restart = {
                "continuous": "continuous",
                "restart_page": "newPage",
                "restart_section": "newSection",
            }[settings.line_numbering]
            line_numbers.set(qn("w:countBy"),
                             str(settings.line_number_count_by))
            line_numbers.set(qn("w:start"),
                             str(settings.line_number_start))
            line_numbers.set(
                qn("w:distance"),
                str(int(settings.line_number_distance_mm * 56.6929)))
            line_numbers.set(qn("w:restart"), restart)
        self._export_hyphenation_settings(settings)
        if settings.header or len(self.docx.sections) > 1:
            header = section.header
            header.is_linked_to_previous = False
            paragraph = header.paragraphs[0]
            paragraph.text = ""
            self._write_header_footer(paragraph, settings.header)
        if settings.footer or len(self.docx.sections) > 1:
            footer = section.footer
            footer.is_linked_to_previous = False
            paragraph = footer.paragraphs[0]
            paragraph.text = ""
            self._write_header_footer(paragraph, settings.footer)
        if self.design.watermark_kind != "none":
            from .design import paint_page_decoration
            from PySide6.QtGui import QPainter

            width, height = mm_to_px(width_mm), mm_to_px(height_mm)
            image = QImage(round(width), round(height), QImage.Format.Format_ARGB32_Premultiplied)
            image.fill(Qt.GlobalColor.transparent)
            painter = QPainter(image)
            watermark = copy.deepcopy(self.design)
            watermark.watermark_layer = "front"
            watermark.page_border = "none"
            paint_page_decoration(painter, QRectF(0, 0, width, height), watermark, "front")
            painter.end()
            section.header.is_linked_to_previous = False
            self._add_anchored_image(section.header.paragraphs[0], image, 0, 0, width, height,
                                     "behind" if self.design.watermark_layer == "behind" else "front",
                                     name="Watermark")

    def _export_hyphenation_settings(self, settings: PageSettings) -> None:
        root = self.docx.settings._element

        def set_boolean(tag: str, enabled: bool) -> None:
            element = root.find(qn(tag))
            if enabled:
                if element is None:
                    element = OxmlElement(tag)
                    root.append(element)
                element.set(qn("w:val"), "1")
            elif element is not None:
                root.remove(element)

        automatic = settings.hyphenation == "automatic"
        set_boolean("w:autoHyphenation", automatic)
        set_boolean("w:doNotHyphenateCaps",
                    automatic and not settings.hyphenate_caps)
        limit = root.find(qn("w:consecutiveHyphenLimit"))
        if automatic and settings.consecutive_hyphen_limit > 0:
            if limit is None:
                limit = OxmlElement("w:consecutiveHyphenLimit")
                root.append(limit)
            limit.set(qn("w:val"),
                      str(settings.consecutive_hyphen_limit))
        elif limit is not None:
            root.remove(limit)

    def _write_header_footer(self, paragraph, text: str) -> None:
        token_re = re.compile(r"\{(page|pages|title|CurrentPage|TotalPages)\}")
        position = 0
        for match in token_re.finditer(text):
            if match.start() > position:
                paragraph.add_run(text[position:match.start()])
            token = match.group(1)
            if token == "title":
                paragraph.add_run(self.state.title)
            elif token in ("page", "CurrentPage"):
                self._add_field(paragraph, "PAGE")
            else:
                self._add_field(paragraph, "NUMPAGES")
            position = match.end()
        if position < len(text):
            paragraph.add_run(text[position:])

    @staticmethod
    def _add_field(paragraph, instruction: str) -> None:
        begin = OxmlElement("w:r")
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        begin.append(fld_begin)
        instr = OxmlElement("w:r")
        instr_text = OxmlElement("w:instrText")
        instr_text.set(qn("xml:space"), "preserve")
        instr_text.text = f" {instruction} "
        instr.append(instr_text)
        end = OxmlElement("w:r")
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        end.append(fld_end)
        paragraph._p.append(begin)
        paragraph._p.append(instr)
        paragraph._p.append(end)

    def _export_body(self) -> None:
        root = self.src.rootFrame()
        iterator = root.begin()
        section_index = 0
        while not iterator.atEnd():
            frame = iterator.currentFrame()
            block = iterator.currentBlock()
            position = (frame.firstPosition() if frame is not None else
                        block.position() if block.isValid() else -1)
            while (section_index + 1 < len(self._sections)
                   and position >= self._sections[section_index + 1][0]):
                section_index += 1
                _start, _end, settings, kind = self._sections[section_index]
                section = self.docx.add_section({
                    "next": WD_SECTION_START.NEW_PAGE,
                    "continuous": WD_SECTION_START.CONTINUOUS,
                    "even": WD_SECTION_START.EVEN_PAGE,
                    "odd": WD_SECTION_START.ODD_PAGE,
                }[kind])
                self._export_sections(section, settings)
            if frame is not None and isinstance(frame, QTextTable):
                self._export_table(frame)
            elif block.isValid():
                paragraph = self.docx.add_paragraph()
                self._source_paragraphs.append((block.position(), paragraph))
                self._export_paragraph_block(block, paragraph)
            iterator += 1

    def _export_floating_objects(self) -> None:
        from .document_features import object_image, paint_objects, visible_objects

        objects = list(visible_objects(self.state.features["objects"]))
        if not objects:
            return
        height = max(1, self.src.pageSize().height())
        paragraphs = {}
        for position, paragraph in self._source_paragraphs:
            block = self.src.findBlock(position)
            page = int(self.src.documentLayout().blockBoundingRect(block).top() / height)
            paragraphs.setdefault(page, paragraph)
        for index, obj in enumerate(objects):
            page = obj.get("page", 0)
            if page not in paragraphs:
                paragraph = self.docx.add_paragraph()
                paragraph.paragraph_format.page_break_before = page > 0
                paragraphs[page] = paragraph
            paragraph = paragraphs[page]
            if obj["kind"] == "image":
                image = object_image(obj)
            else:
                scale = min(1.0, 4096 / max(obj["width"], obj["height"]))
                image = QImage(max(1, round(obj["width"] * scale)),
                               max(1, round(obj["height"] * scale)),
                               QImage.Format.Format_ARGB32_Premultiplied)
                image.fill(Qt.GlobalColor.transparent)
                from PySide6.QtGui import QPainter

                painter = QPainter(image)
                painter.scale(scale, scale)
                local = dict(obj, x=0, y=0, group="")
                paint_objects(painter, [local], page,
                              "behind" if obj.get("wrap") == "behind" else "front")
                painter.end()
            if not image.isNull():
                self._add_anchored_image(paragraph, image, obj["x"], obj["y"],
                                         obj["width"], obj["height"], obj.get("wrap", "front"),
                                         index + 1, obj.get("name", "Floating object"))

    @staticmethod
    def _add_anchored_image(paragraph, image, x, y, width, height, wrap, z=1, name="Object"):
        from PySide6.QtCore import QBuffer, QByteArray, QIODevice

        data = QByteArray()
        buffer = QBuffer(data)
        buffer.open(QIODevice.OpenModeFlag.WriteOnly)
        image.save(buffer, "PNG")
        buffer.close()
        run = paragraph.add_run()
        shape = run.add_picture(BytesIO(bytes(data.data())),
                                width=Emu(round(width * EMU_PER_PX)),
                                height=Emu(round(height * EMU_PER_PX)))
        inline = shape._inline
        anchor = OxmlElement("wp:anchor")
        for key, value in {"distT": "0", "distB": "0", "distL": "0", "distR": "0",
                           "simplePos": "0", "relativeHeight": str(z),
                           "behindDoc": "1" if wrap == "behind" else "0",
                           "locked": "0", "layoutInCell": "1", "allowOverlap": "1"}.items():
            anchor.set(key, value)
        simple = OxmlElement("wp:simplePos")
        simple.set("x", "0")
        simple.set("y", "0")
        anchor.append(simple)
        for axis, offset in (("H", x), ("V", y)):
            position = OxmlElement("wp:position" + axis)
            position.set("relativeFrom", "page")
            value = OxmlElement("wp:posOffset")
            value.text = str(round(offset * EMU_PER_PX))
            position.append(value)
            anchor.append(position)
        anchor.append(inline.find(qn("wp:extent")))
        if wrap in ("square", "tight"):
            wrapping = OxmlElement("wp:wrapSquare" if wrap == "square" else "wp:wrapTight")
            wrapping.set("wrapText", "bothSides")
            if wrap == "tight":
                polygon = OxmlElement("wp:wrapPolygon")
                polygon.set("edited", "0")
                for index, (px, py) in enumerate(((0, 0), (21600, 0), (21600, 21600),
                                                 (0, 21600), (0, 0))):
                    point = OxmlElement("wp:start" if index == 0 else "wp:lineTo")
                    point.set("x", str(px))
                    point.set("y", str(py))
                    polygon.append(point)
                wrapping.append(polygon)
            anchor.append(wrapping)
        else:
            anchor.append(OxmlElement("wp:wrapNone"))
        properties = inline.find(qn("wp:docPr"))
        properties.set("name", name)
        anchor.append(properties)
        graphic_properties = inline.find(qn("wp:cNvGraphicFramePr"))
        if graphic_properties is not None:
            anchor.append(graphic_properties)
        anchor.append(inline.find(qn("a:graphic")))
        inline.getparent().replace(inline, anchor)

    def _export_paragraph_block(self, block, paragraph) -> None:
        fmt = block.blockFormat()
        break_kind = block_break_kind(block)
        preset = _preset_for_block(block)
        if preset:
            try:
                paragraph.style = self.docx.styles[preset]
            except KeyError:
                pass
        pf = paragraph.paragraph_format
        alignment = fmt.alignment() & Qt.AlignmentFlag.AlignHorizontal_Mask
        if alignment in _ALIGNMENT_MAP:
            pf.alignment = _ALIGNMENT_MAP[alignment]
        if fmt.topMargin() > 0:
            pf.space_before = Pt(_px_to_pt(fmt.topMargin()))
        if fmt.bottomMargin() > 0:
            pf.space_after = Pt(_px_to_pt(fmt.bottomMargin()))
        if fmt.leftMargin() > 0:
            pf.left_indent = Pt(_px_to_pt(fmt.leftMargin()))
        if fmt.rightMargin() > 0:
            pf.right_indent = Pt(_px_to_pt(fmt.rightMargin()))
        if fmt.textIndent() != 0:
            pf.first_line_indent = Pt(_px_to_pt(fmt.textIndent()))
        if fmt.lineHeightType() == QTextBlockFormat.LineHeightTypes.ProportionalHeight.value:
            pf.line_spacing = fmt.lineHeight() / 100.0
        elif fmt.lineHeightType() in (QTextBlockFormat.LineHeightTypes.FixedHeight.value,
                                     QTextBlockFormat.LineHeightTypes.MinimumHeight.value):
            pf.line_spacing = Pt(_px_to_pt(fmt.lineHeight()))
            pf.line_spacing_rule = (WD_LINE_SPACING.EXACTLY
                                    if fmt.lineHeightType() == QTextBlockFormat.LineHeightTypes.FixedHeight.value
                                    else WD_LINE_SPACING.AT_LEAST)
        from .formatting import paragraph_border

        border = paragraph_border(block)
        if border:
            container = OxmlElement("w:pBdr")
            for side in border["sides"]:
                element = OxmlElement("w:" + side)
                element.set(qn("w:val"), "single")
                element.set(qn("w:color"), border["color"][1:])
                element.set(qn("w:sz"), str(max(2, round(border["width"] * PT_PER_PX * 8))))
                element.set(qn("w:space"), "0")
                container.append(element)
            paragraph._p.get_or_add_pPr().append(container)
        if fmt.background().style() != Qt.BrushStyle.NoBrush:
            shading = OxmlElement("w:shd")
            shading.set(qn("w:val"), "clear")
            shading.set(qn("w:fill"), fmt.background().color().name()[1:])
            paragraph._p.get_or_add_pPr().append(shading)
        if break_kind == "page":
            pf.page_break_before = True
        elif break_kind == "column":
            paragraph.add_run().add_break(WD_BREAK.COLUMN)
        text_list = block.textList()
        if text_list is not None:
            num_id = self.numbering.num_id(text_list)
            ilvl = max(0, text_list.format().indent() - 1)
            p_pr = paragraph._p.get_or_add_pPr()
            num_pr = OxmlElement("w:numPr")
            ilvl_el = OxmlElement("w:ilvl")
            ilvl_el.set(qn("w:val"), str(min(ilvl, 8)))
            num_id_el = OxmlElement("w:numId")
            num_id_el.set(qn("w:val"), str(num_id))
            num_pr.append(ilvl_el)
            num_pr.append(num_id_el)
            p_pr.append(num_pr)
        iterator = block.begin()
        while not iterator.atEnd():
            fragment = iterator.fragment()
            if fragment.isValid():
                self._export_fragment(fragment, paragraph)
            iterator += 1

    def _export_fragment(self, fragment, paragraph) -> None:
        from .document_features import ANCHOR

        keys = [name[len(ANCHOR):] for name in fragment.charFormat().anchorNames()
                if name.startswith(ANCHOR) and name[len(ANCHOR):] in self._ranges
                and self._ranges[name[len(ANCHOR):]][0] <= fragment.position()
                < self._ranges[name[len(ANCHOR):]][1]]
        for key in keys:
            start, _end = self._ranges[key]
            if start == fragment.position():
                bookmark = OxmlElement("w:bookmarkStart")
                identifier, name = self._bookmarks[key]
                bookmark.set(qn("w:id"), str(identifier))
                bookmark.set(qn("w:name"), name)
                paragraph._p.append(bookmark)
        node_key = next((key for key in keys
                         if self.state.features["nodes"].get(key, {}).get("kind")
                         in ("reference", "datetime", "signature", "attachment")
                         or key in self._footnotes), None)
        if node_key:
            if node_key not in self._emitted_fields:
                self._export_node(node_key, fragment, paragraph)
                self._emitted_fields.add(node_key)
        else:
            self._export_fragment_content(fragment, paragraph)
        for key in reversed(keys):
            _start, end = self._ranges[key]
            if end == fragment.position() + fragment.length():
                bookmark = OxmlElement("w:bookmarkEnd")
                bookmark.set(qn("w:id"), str(self._bookmarks[key][0]))
                paragraph._p.append(bookmark)

    def _export_node(self, key, fragment, paragraph) -> None:
        node = self.state.features["nodes"][key]
        kind = node["kind"]
        start, end = self._ranges[key]
        cursor = QTextCursor(self.src)
        cursor.setPosition(start)
        cursor.setPosition(end, QTextCursor.MoveMode.KeepAnchor)
        text = cursor.selectedText().replace("\u2029", " ").replace("\u2063", "")
        if key in self._footnotes:
            run = paragraph.add_run()
            run.font.superscript = True
            marker = OxmlElement("w:footnoteReference")
            marker.set(qn("w:id"), str(self._footnotes[key]))
            run._r.append(marker)
        elif kind in ("datetime", "reference"):
            if kind == "datetime":
                style = node.get("format", "date")
                pattern = {"date": "MMMM dd, yyyy", "time": "HH:mm",
                           "datetime": "yyyy-MM-dd HH:mm"}.get(style, "yyyy-MM-dd")
                instruction = ('TIME' if style == "time" else 'DATE') + ' \\@ "' + pattern + '"'
            else:
                target = self._bookmarks.get(node.get("target"))
                if target is None:
                    paragraph.add_run("[Reference missing]")
                    return
                instruction = ("PAGEREF" if node.get("display") == "page" else "REF")
                instruction += " " + target[1] + " \\h"
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), instruction)
            field.set(qn("w:dirty"), "true")
            run = paragraph.add_run(text)
            self._apply_char_format(run, fragment.charFormat())
            field.append(run._r)
            paragraph._p.append(field)
        elif kind == "signature":
            container = OxmlElement("w:sdt")
            properties = OxmlElement("w:sdtPr")
            tag = OxmlElement("w:tag")
            tag.set(qn("w:val"), "FolioSignature:" + key)
            alias = OxmlElement("w:alias")
            alias.set(qn("w:val"), "Signature")
            properties.extend((tag, alias, OxmlElement("w:text")))
            container.append(properties)
            content = OxmlElement("w:sdtContent")
            run = paragraph.add_run(text or node.get("signed_by") or "Signature: __________________")
            self._apply_char_format(run, fragment.charFormat())
            content.append(run._r)
            container.append(content)
            paragraph._p.append(container)
        elif kind == "attachment":
            suffix = Path(node.get("name", "attachment.bin")).suffix.lower()
            if not re.fullmatch(r"\.[a-z0-9]{1,10}", suffix):
                suffix = ".bin"
            content_type = {
                ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ".pdf": "application/pdf",
            }.get(suffix, "application/octet-stream")
            name = PackURI("/word/embeddings/Folio" + str(len(self._emitted_fields)) + suffix)
            embedded = Part(name, content_type, base64.b64decode(node["data"], validate=True),
                            self.docx.part.package)
            paragraph.part.relate_to(embedded, RT.PACKAGE)
            run = paragraph.add_run(text or "Attachment: " + node.get("name", "file"))
            self._apply_char_format(run, fragment.charFormat())

    def _export_fragment_content(self, fragment, paragraph) -> None:
        fmt = fragment.charFormat()
        column_marker = (fmt.isAnchor()
                         and fmt.anchorHref() == COLUMN_BREAK_URL
                         and COLUMN_BREAK_MARKER in fragment.text())
        if fmt.isImageFormat():
            image_fmt = QTextImageFormat(fmt.toImageFormat())
            image = _decode_image_name(image_fmt.name())
            if image is None:
                return
            buffer = BytesIO()
            from PySide6.QtCore import QBuffer, QByteArray, QIODevice

            data = QByteArray()
            qbuffer = QBuffer(data)
            qbuffer.open(QIODevice.OpenModeFlag.WriteOnly)
            image.save(qbuffer, "PNG")
            qbuffer.close()
            buffer.write(bytes(data.data()))
            buffer.seek(0)
            run = paragraph.add_run()
            width = image_fmt.width() or image.width()
            height = image_fmt.height() or image.height()
            picture = run.add_picture(
                buffer,
                width=Emu(int(width * EMU_PER_PX)),
                height=Emu(int(height * EMU_PER_PX)),
            )
            description = fmt.property(QTextFormat.Property.ImageAltText)
            if description:
                picture._inline.docPr.set("descr", str(description))
            return
        text = fragment.text()
        from .document_features import ANCHOR, MARKER

        if any(name.startswith(ANCHOR) for name in fmt.anchorNames()):
            text = text.replace(MARKER, "")
        if column_marker:
            text = text.replace(COLUMN_BREAK_MARKER, "")
        if not text:
            return
        if fmt.isAnchor() and fmt.anchorHref() and not column_marker:
            self._export_hyperlink(paragraph, fragment, fmt, text)
            return
        lines = text.split(" ")
        run = paragraph.add_run(lines[0])
        for line in lines[1:]:
            run.add_break()
            run.add_text(line)
        self._apply_char_format(run, fmt)

    def _export_hyperlink(self, paragraph, fragment, fmt, text: str) -> None:
        part = paragraph.part
        rel_id = part.relate_to(fmt.anchorHref(), RT.HYPERLINK, is_external=True)
        run = paragraph.add_run(text)
        self._apply_char_format(run, fmt)
        element = run._r
        element.getparent().remove(element)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), rel_id)
        hyperlink.append(element)
        paragraph._p.append(hyperlink)

    @staticmethod
    def _apply_char_format(run, fmt: QTextCharFormat) -> None:
        font = run.font
        if fmt.fontPointSize() > 0:
            font.size = Pt(fmt.fontPointSize())
        family = fmt.font().family()
        if family:
            font.name = family.split(",")[0]
        font.bold = fmt.fontWeight() >= QFont.Weight.Bold
        font.italic = bool(fmt.fontItalic())
        if fmt.fontUnderline() or fmt.underlineStyle() not in (
            QTextCharFormat.UnderlineStyle.NoUnderline,
        ):
            font.underline = True
        if fmt.fontStrikeOut():
            font.strike = True
        alignment = fmt.verticalAlignment()
        if alignment == QTextCharFormat.VerticalAlignment.AlignSubScript:
            font.subscript = True
        elif alignment == QTextCharFormat.VerticalAlignment.AlignSuperScript:
            font.superscript = True
        if fmt.textOutline().style() != Qt.PenStyle.NoPen:
            properties = run._r.get_or_add_rPr()
            shadow = OxmlElement("w:shadow")
            shadow.set(qn("w:val"), "1")
            properties.append(shadow)
        if fmt.hasProperty(QTextFormat.Property.ForegroundBrush):
            color = fmt.foreground().color()
            font.color.rgb = RGBColor(color.red(), color.green(), color.blue())
        if (fmt.hasProperty(QTextFormat.Property.BackgroundBrush)
                and fmt.background().style() != Qt.BrushStyle.NoBrush):
            background = fmt.background().color()
            index = _nearest_highlight(background)
            if index is not None:
                font.highlight_color = index


def _export_table_impl(exporter: "_Exporter", table: QTextTable) -> None:
    rows = table.rows()
    cols = table.columns()
    docx_table = exporter.docx.add_table(rows=rows, cols=cols)
    try:
        docx_table.style = exporter.docx.styles["Table Grid"]
    except KeyError:
        pass
    fmt = table.format()
    widths = fmt.columnWidthConstraints()
    if len(widths) == cols:
        docx_table.autofit = False
        available = mm_to_px(exporter.state.page.column_width_mm())
        for col, constraint in enumerate(widths):
            if constraint.type() == QTextLength.Type.VariableLength:
                continue
            width = constraint.value(available)
            docx_table.columns[col].width = Emu(round(width * EMU_PER_PX))
            for row in range(rows):
                docx_table.cell(row, col).width = Emu(round(width * EMU_PER_PX))
    for row in range(rows):
        heights = []
        for col in range(cols):
            block = table.cellAt(row, col).firstCursorPosition().block()
            block_format = block.blockFormat()
            if block_format.lineHeightType() == QTextBlockFormat.LineHeightTypes.MinimumHeight.value:
                heights.append(block_format.lineHeight())
        if heights:
            docx_table.rows[row].height = Emu(round(max(heights) * EMU_PER_PX))
            docx_table.rows[row].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    tbl_pr = docx_table._tbl.tblPr
    cell_mar = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(fmt.cellPadding() * 20 * PT_PER_PX)))
        el.set(qn("w:type"), "dxa")
        cell_mar.append(el)
    tbl_pr.append(cell_mar)
    for row in range(rows):
        for col in range(cols):
            cell = table.cellAt(row, col)
            if cell.row() != row or cell.column() != col:
                continue
            target = docx_table.cell(row, col)
            span = (cell.rowSpan(), cell.columnSpan())
            if span != (1, 1):
                try:
                    target = target.merge(
                        docx_table.cell(row + span[0] - 1, col + span[1] - 1)
                    )
                except Exception:
                    target = docx_table.cell(row, col)
            _fill_docx_cell(exporter, cell, target)
            if row == 0 and fmt.headerRowCount() > 0:
                tc_pr = target._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                background = cell.format().background().color().name()
                shd.set(qn("w:fill"), background.lstrip("#"))
                tc_pr.append(shd)


def _fill_docx_cell(exporter: "_Exporter", cell, docx_cell) -> None:
    first = cell.firstPosition()
    last = cell.lastPosition()
    doc = exporter.src
    block = doc.findBlock(first)
    first_block = True
    while block.isValid() and block.position() < last:
        if first_block:
            paragraph = docx_cell.paragraphs[0]
            for run in list(paragraph.runs):
                run._r.getparent().remove(run._r)
            first_block = False
        else:
            paragraph = docx_cell.add_paragraph()
        exporter._export_paragraph_block(block, paragraph)
        block = block.next()


_Exporter._export_table = _export_table_impl


def _collect_members(docx_bytes: bytes) -> dict[str, bytes]:
    members: dict[str, bytes] = {}
    with zipfile.ZipFile(BytesIO(docx_bytes)) as zin:
        for info in zin.infolist():
            members[info.filename] = zin.read(info.filename)
    return members


def _write_members(members: dict[str, bytes]) -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in members.items():
            zout.writestr(name, data)
    return buffer.getvalue()


def _embed_state(docx_bytes: bytes, state: DocumentState,
                 numbering_xml: bytes | None) -> bytes:
    members = _collect_members(docx_bytes)
    if numbering_xml is not None:
        members["word/numbering.xml"] = numbering_xml
        content_types = members["[Content_Types].xml"].decode("utf-8")
        if "numbering+xml" not in content_types:
            content_types = content_types.replace(
                "</Types>",
                '<Override PartName="/word/numbering.xml" ContentType='
                '"application/vnd.openxmlformats-officedocument.'
                'wordprocessingml.numbering+xml"/></Types>',
            )
            members["[Content_Types].xml"] = content_types.encode("utf-8")
        rels_name = "word/_rels/document.xml.rels"
        rels = members[rels_name].decode("utf-8")
        if "relationships/numbering" not in rels:
            used = set(re.findall(r'Id="(rId\d+)"', rels))
            new_id = "rId9000"
            counter = 9000
            while new_id in used:
                counter += 1
                new_id = f"rId{counter}"
            rels = rels.replace(
                "</Relationships>",
                f'<Relationship Id="{new_id}" Type='
                '"http://schemas.openxmlformats.org/officeDocument/2006/'
                'relationships/numbering" Target="numbering.xml"/>'
                '</Relationships>',
            )
            members[rels_name] = rels.encode("utf-8")
    digest = _word_digest(members)
    metadata = json.dumps(
        {
            "schema_version": 1,
            "word_digest": digest,
            "state": state.to_dict(),
        },
        ensure_ascii=False,
    ).encode("utf-8")
    if len(metadata) > MAX_METADATA_BYTES:
        raise DocxError(
            "The embedded Folio metadata exceeds the "
            f"{MAX_METADATA_BYTES // (1024 * 1024)} MiB limit; the file "
            "was not written. Reduce embedded images or comments."
        )
    content_types = members["[Content_Types].xml"].decode("utf-8")
    if "/customXml/folio-state.json" not in content_types:
        content_types = content_types.replace(
            "</Types>",
            f'<Override PartName="/{META_PART}" ContentType='
            f'"{META_CONTENT_TYPE}"/></Types>',
        )
    members["[Content_Types].xml"] = content_types.encode("utf-8")
    rels = members["_rels/.rels"].decode("utf-8")
    used = set(re.findall(r'Id="(rId\d+)"', rels))
    new_id = "rId9000"
    counter = 9000
    while new_id in used:
        counter += 1
        new_id = f"rId{counter}"
    rels = rels.replace(
        "</Relationships>",
        f'<Relationship Id="{new_id}" Type="{META_REL_TYPE}" '
        f'Target="{META_PART}"/></Relationships>',
    )
    members["_rels/.rels"] = rels.encode("utf-8")
    members[META_PART] = metadata
    return _write_members(members)


def write_docx(path: Path, state: DocumentState, document: QTextDocument) -> None:
    snapshot = copy.deepcopy(state)
    snapshot.html = document.toHtml()
    try:
        snapshot.page.validate()
        from .design import DesignSettings
        from .document_features import checked_features

        snapshot.design = DesignSettings.from_dict(snapshot.design).to_dict()
        snapshot.features = checked_features(snapshot.features)
    except Exception as exc:
        raise DocxError(f"Document settings are invalid: {exc}") from exc
    exporter = _Exporter(snapshot, document)
    raw = exporter.build()
    final = _embed_state(raw, snapshot, exporter.numbering.numbering_xml())
    atomic_write(Path(path), final)


def _preflight_zip(data: bytes) -> zipfile.ZipFile:
    if not zipfile.is_zipfile(BytesIO(data)):
        raise DocxError("This file is not a valid DOCX (not a ZIP archive).")
    try:
        archive = zipfile.ZipFile(BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise DocxError(f"The DOCX archive is corrupt: {exc}") from exc
    infos = archive.infolist()
    if len(infos) > MAX_ZIP_ENTRIES:
        archive.close()
        raise DocxError("This DOCX has too many entries to open safely.")
    total = sum(info.file_size for info in infos)
    if total > MAX_UNCOMPRESSED_BYTES:
        archive.close()
        raise DocxError("This DOCX is too large to open safely.")
    return archive


def _read_embedded_state(archive: zipfile.ZipFile) -> dict | None:
    names = set(archive.namelist())
    if "_rels/.rels" not in names or META_PART not in names:
        return None
    try:
        rels = archive.read("_rels/.rels").decode("utf-8", errors="replace")
    except Exception:
        return None
    if META_REL_TYPE not in rels:
        return None
    try:
        return json.loads(archive.read(META_PART).decode("utf-8"))
    except Exception:
        return None


def _read_hf_text(container) -> str:
    pieces: list[str] = []
    for paragraph in container.paragraphs:
        for child in paragraph._p.iter():
            tag = child.tag
            if tag == qn("w:t"):
                pieces.append(child.text or "")
            elif tag == qn("w:instrText"):
                instruction = (child.text or "").strip()
                if instruction == "PAGE":
                    pieces.append("{page}")
                elif instruction == "NUMPAGES":
                    pieces.append("{pages}")
        pieces.append("\n")
    return "".join(pieces).strip()


def _numbering_map(docx_doc) -> dict[tuple[str, int], str]:
    try:
        part = docx_doc.part.numbering_part
    except Exception:
        return {}
    element = part.element
    abstract_fmts: dict[tuple[str, int], str] = {}
    for abstract in element.findall(qn("w:abstractNum")):
        aid = abstract.get(qn("w:abstractNumId"))
        for lvl in abstract.findall(qn("w:lvl")):
            try:
                ilvl = int(lvl.get(qn("w:ilvl"), "0"))
            except ValueError:
                ilvl = 0
            fmt_el = lvl.find(qn("w:numFmt"))
            abstract_fmts[(aid, ilvl)] = (
                fmt_el.get(qn("w:val")) if fmt_el is not None else "decimal"
            )
    result: dict[tuple[str, int], str] = {}
    for num in element.findall(qn("w:num")):
        nid = num.get(qn("w:numId"))
        aid_el = num.find(qn("w:abstractNumId"))
        aid = aid_el.get(qn("w:val")) if aid_el is not None else None
        for ilvl in range(9):
            result[(nid, ilvl)] = abstract_fmts.get((aid, ilvl), "decimal")
    return result


def _paragraph_num_pr(paragraph):
    p_pr = paragraph._p.find(qn("w:pPr"))
    if p_pr is not None:
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is not None:
            return num_pr
    style = paragraph.style
    if style is not None:
        p_pr = style._element.find(qn("w:pPr"))
        if p_pr is not None:
            return p_pr.find(qn("w:numPr"))
    return None


def _import_paragraph(cursor, paragraph, pending_break: bool,
                      list_map: dict, num_map: dict,
                      warnings: list[str]) -> bool:
    block_fmt = QTextBlockFormat()
    style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
    heading_match = re.match(r"Heading (\d)", style_name)
    if heading_match:
        block_fmt.setHeadingLevel(int(heading_match.group(1)))
    if pending_break or paragraph.paragraph_format.page_break_before:
        block_fmt.setPageBreakPolicy(QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore)
    alignment = paragraph.alignment
    align_map = {
        WD_ALIGN_PARAGRAPH.LEFT: Qt.AlignmentFlag.AlignLeft,
        WD_ALIGN_PARAGRAPH.RIGHT: Qt.AlignmentFlag.AlignRight,
        WD_ALIGN_PARAGRAPH.CENTER: Qt.AlignmentFlag.AlignHCenter,
        WD_ALIGN_PARAGRAPH.JUSTIFY: Qt.AlignmentFlag.AlignJustify,
    }
    if alignment in align_map:
        block_fmt.setAlignment(align_map[alignment])
    pf = paragraph.paragraph_format
    if pf.space_before is not None:
        block_fmt.setTopMargin(pf.space_before.pt / PT_PER_PX)
    if pf.space_after is not None:
        block_fmt.setBottomMargin(pf.space_after.pt / PT_PER_PX)
    if pf.left_indent is not None:
        block_fmt.setLeftMargin(pf.left_indent.pt / PT_PER_PX)
    if pf.right_indent is not None:
        block_fmt.setRightMargin(pf.right_indent.pt / PT_PER_PX)
    if pf.first_line_indent is not None:
        block_fmt.setTextIndent(pf.first_line_indent.pt / PT_PER_PX)
    if pf.line_spacing is not None and isinstance(pf.line_spacing, float):
        block_fmt.setLineHeight(
            pf.line_spacing * 100.0,
            QTextBlockFormat.LineHeightTypes.ProportionalHeight.value,
        )
    p_pr = paragraph._p.find(qn("w:pPr"))
    if p_pr is not None:
        shading = p_pr.find(qn("w:shd"))
        if shading is not None:
            fill = shading.get(qn("w:fill"), "")
            if re.fullmatch(r"[0-9a-fA-F]{6}", fill):
                block_fmt.setBackground(QColor("#" + fill))
        borders = p_pr.find(qn("w:pBdr"))
        if borders is not None:
            from .formatting import PARAGRAPH_BORDER_PROPERTY

            sides = [side for side in ("top", "bottom", "left", "right")
                     if borders.find(qn("w:" + side)) is not None
                     and borders.find(qn("w:" + side)).get(qn("w:val")) not in ("none", "nil")]
            if sides:
                border = borders.find(qn("w:" + sides[0]))
                color = border.get(qn("w:color"), "000000")
                color = color if re.fullmatch(r"[0-9a-fA-F]{6}", color) else "000000"
                try:
                    width = max(0.25, min(12, int(border.get(qn("w:sz"), "8")) / 8 / PT_PER_PX))
                except ValueError:
                    width = 1.0
                block_fmt.setProperty(PARAGRAPH_BORDER_PROPERTY, json.dumps({
                    "sides": sides, "color": "#" + color, "width": width}))
    cursor.setBlockFormat(block_fmt)
    num_pr = _paragraph_num_pr(paragraph)
    if num_pr is not None:
        num_id_el = num_pr.find(qn("w:numId"))
        ilvl_el = num_pr.find(qn("w:ilvl"))
        num_id = num_id_el.get(qn("w:val")) if num_id_el is not None else "0"
        try:
            ilvl = int(ilvl_el.get(qn("w:val"))) if ilvl_el is not None else 0
        except ValueError:
            ilvl = 0
        num_fmt = num_map.get((num_id, ilvl), "decimal")
        key = (num_id, ilvl)
        existing = list_map.get(key)
        if existing is None:
            fmt = QTextListFormat()
            fmt.setStyle({
                "bullet": QTextListFormat.Style.ListDisc,
                "lowerLetter": QTextListFormat.Style.ListLowerAlpha,
                "upperLetter": QTextListFormat.Style.ListUpperAlpha,
                "lowerRoman": QTextListFormat.Style.ListLowerRoman,
                "upperRoman": QTextListFormat.Style.ListUpperRoman,
            }.get(num_fmt, QTextListFormat.Style.ListDecimal))
            fmt.setIndent(ilvl + 1)
            existing = cursor.createList(fmt)
            list_map[key] = existing
        else:
            existing.add(cursor.block())
    next_break = False
    for child in paragraph._p:
        if child.tag == qn("w:hyperlink"):
            rel_id = child.get(qn("r:id"))
            target = None
            if rel_id:
                try:
                    target = paragraph.part.rels[rel_id].target_ref
                except Exception:
                    target = None
            for run_el in child.findall(qn("w:r")):
                fmt = QTextCharFormat()
                if target:
                    fmt.setAnchor(True)
                    fmt.setAnchorHref(target)
                fmt.setForeground(QColor("#185abd"))
                fmt.setFontUnderline(True)
                for part in run_el:
                    if part.tag == qn("w:t") and part.text:
                        cursor.insertText(part.text, fmt)
                    elif part.tag == qn("w:tab"):
                        cursor.insertText("\t", fmt)
        elif child.tag == qn("w:r"):
            next_break = _import_run(cursor, child, paragraph) or next_break
        elif child.tag in (qn("w:fldSimple"), qn("w:sdt")):
            for run_el in child.iter(qn("w:r")):
                next_break = _import_run(cursor, run_el, paragraph) or next_break
    if style_name == "Title":
        _retitle_block(cursor, 32, "#17365d")
    elif style_name == "Subtitle":
        _retitle_block(cursor, 16, "#64748b")
    return next_break


def _retitle_block(cursor, size: float, color: str) -> None:
    edit = QTextCursor(cursor.block())
    edit.movePosition(QTextCursor.MoveOperation.StartOfBlock)
    edit.movePosition(QTextCursor.MoveOperation.EndOfBlock,
                      QTextCursor.MoveMode.KeepAnchor)
    fmt = QTextCharFormat()
    fmt.setFontPointSize(size)
    fmt.setForeground(QColor(color))
    edit.mergeCharFormat(fmt)


def _style_font_chain(paragraph):
    fonts = []
    style = paragraph.style
    while style is not None:
        fonts.append(style.font)
        style = style.base_style
    return fonts


def _run_char_format(run, paragraph) -> QTextCharFormat:
    fmt = QTextCharFormat()
    if run.bold:
        fmt.setFontWeight(QFont.Weight.Bold)
    elif run.bold is False:
        fmt.setFontWeight(QFont.Weight.Normal)
    if run.italic:
        fmt.setFontItalic(True)
    if run.underline:
        fmt.setFontUnderline(True)
    if run.font.strike:
        fmt.setFontStrikeOut(True)
    if run.font.subscript:
        fmt.setVerticalAlignment(QTextCharFormat.VerticalAlignment.AlignSubScript)
    elif run.font.superscript:
        fmt.setVerticalAlignment(QTextCharFormat.VerticalAlignment.AlignSuperScript)
    run_props = run._r.find(qn("w:rPr"))
    if run_props is not None and run_props.find(qn("w:shadow")) is not None:
        fmt.setTextOutline(QPen(QColor("#78869a"), 0.65))
    if run.font.size is not None:
        fmt.setFontPointSize(run.font.size.pt)
    else:
        for font in _style_font_chain(paragraph):
            if font.size is not None:
                fmt.setFontPointSize(font.size.pt)
                break
    if run.font.name:
        fmt.setFontFamilies([run.font.name])
    else:
        for font in _style_font_chain(paragraph):
            if font.name:
                fmt.setFontFamilies([font.name])
                break
    try:
        if run.font.color is not None and run.font.color.rgb is not None:
            rgb = run.font.color.rgb
            fmt.setForeground(QColor(rgb[0], rgb[1], rgb[2]))
    except Exception:
        pass
    try:
        if run.font.highlight_color is not None:
            highlight_map = {
                WD_COLOR_INDEX.YELLOW: "#ffff00",
                WD_COLOR_INDEX.BRIGHT_GREEN: "#00ff00",
                WD_COLOR_INDEX.TURQUOISE: "#00ffff",
                WD_COLOR_INDEX.PINK: "#ff00ff",
                WD_COLOR_INDEX.RED: "#ff0000",
                WD_COLOR_INDEX.BLUE: "#0000ff",
            }
            hex_color = highlight_map.get(run.font.highlight_color)
            if hex_color:
                fmt.setBackground(QColor(hex_color))
    except Exception:
        pass
    return fmt


def _import_drawing(cursor, drawing_el, paragraph) -> None:
    blip = None
    for blip in drawing_el.iter(qn("a:blip")):
        break
    if blip is None:
        return
    embed = blip.get(qn("r:embed"))
    if not embed:
        return
    try:
        image_part = paragraph.part.related_parts[embed]
        image = QImage.fromData(image_part.blob)
    except Exception:
        image = QImage()
    if image.isNull():
        return
    from .editor import _image_to_data_uri

    fmt = QTextImageFormat()
    fmt.setName(_image_to_data_uri(image))
    extent = drawing_el.find(".//" + qn("wp:extent"))
    width_px = image.width()
    if extent is not None and extent.get("cx"):
        try:
            width_px = int(extent.get("cx")) / EMU_PER_PX
        except ValueError:
            width_px = image.width()
    fmt.setWidth(min(float(width_px), 640.0))
    fmt.setHeight(fmt.width() * image.height() / max(1, image.width()))
    properties = drawing_el.find(".//" + qn("wp:docPr"))
    if properties is not None and properties.get("descr"):
        fmt.setProperty(QTextFormat.Property.ImageAltText, properties.get("descr"))
    cursor.insertImage(fmt)


def _import_run(cursor, run_el, paragraph) -> bool:
    from docx.text.run import Run

    run = Run(run_el, paragraph)
    fmt = _run_char_format(run, paragraph)
    for child in run_el:
        tag = child.tag
        if tag == qn("w:t"):
            if child.text:
                cursor.insertText(child.text, fmt)
        elif tag == qn("w:tab"):
            cursor.insertText("\t", fmt)
        elif tag in (qn("w:footnoteReference"), qn("w:endnoteReference")):
            _import_note(cursor, child, paragraph)
        elif tag in (qn("w:br"), qn("w:cr")):
            break_type = child.get(qn("w:type")) if tag == qn("w:br") else ""
            if break_type in ("page", "column"):
                cursor.insertBlock()
                block_fmt = cursor.blockFormat()
                block_fmt.setPageBreakPolicy(
                    QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore)
                cursor.setBlockFormat(block_fmt)
                if break_type == "column":
                    marker = QTextCharFormat()
                    marker.setAnchor(True)
                    marker.setAnchorHref(COLUMN_BREAK_URL)
                    marker.setFontPointSize(0.1)
                    marker.setForeground(QColor(Qt.GlobalColor.transparent))
                    cursor.insertText(COLUMN_BREAK_MARKER, marker)
            else:
                cursor.insertText(" ", fmt)
        elif tag == qn("w:drawing"):
            _import_drawing(cursor, child, paragraph)
    return False


def _import_note(cursor, reference, paragraph) -> None:
    from docx.oxml import parse_xml
    from .document_features import ANCHOR, uid

    kind = "footnote" if reference.tag == qn("w:footnoteReference") else "endnote"
    relation = RT.FOOTNOTES if kind == "footnote" else RT.ENDNOTES
    state = getattr(cursor.document(), "folio_state", None)
    if state is None:
        return
    try:
        part = paragraph.part.part_related_by(relation)
        root = parse_xml(part.blob)
        identifier = reference.get(qn("w:id"))
        note = next(item for item in root if item.get(qn("w:id")) == identifier)
        lines = ["".join(text.text or "" for text in item.iter(qn("w:t")))
                 for item in note if item.tag == qn("w:p")]
        text = "\n".join(lines).strip()
    except (KeyError, StopIteration, ValueError):
        cursor.insertText(f"[Missing {kind}]")
        return
    key = uid()
    number = 1 + sum(node.get("kind") == "note" and node.get("type") == kind
                     for node in state.features["nodes"].values())
    state.features["nodes"][key] = {"kind": "note", "type": kind, "text": text or "(empty note)"}
    fmt = QTextCharFormat()
    fmt.setAnchor(True)
    fmt.setAnchorNames([ANCHOR + key])
    fmt.setAnchorHref("folio-note:" + key)
    fmt.setVerticalAlignment(QTextCharFormat.VerticalAlignment.AlignSuperScript)
    cursor.insertText(str(number), fmt)


def _import_body(data: bytes, source_name: str) -> tuple[DocumentState, list[str]]:
    warnings = [IMPORT_WARNING]
    try:
        docx_doc = docx.Document(BytesIO(data))
    except Exception as exc:
        raise DocxError(f"Could not read the DOCX body: {exc}") from exc
    document = SafeDocument()
    state = DocumentState()
    document.folio_state = state
    document.setDefaultFont(QFont("Calibri", 11))
    cursor = QTextCursor(document)
    list_map: dict = {}
    num_map = _numbering_map(docx_doc)
    pending_break = False
    needs_block = False
    for element in docx_doc.iter_inner_content():
        if isinstance(element, DocxTable):
            pending_break = False
            cursor.movePosition(QTextCursor.MoveOperation.End)
            if needs_block:
                cursor.insertBlock()
            _import_table(cursor, element, warnings)
            cursor.movePosition(QTextCursor.MoveOperation.End)
            needs_block = bool(cursor.block().text())
        elif isinstance(element, DocxParagraph):
            cursor.movePosition(QTextCursor.MoveOperation.End)
            if needs_block:
                cursor.insertBlock()
            pending_break = _import_paragraph(
                cursor, element, pending_break, list_map, num_map, warnings)
            needs_block = True
    endnotes = [node for node in state.features["nodes"].values()
                if node.get("kind") == "note" and node.get("type") == "endnote"]
    if endnotes:
        from .document_features import ANCHOR, uid
        cursor.movePosition(QTextCursor.MoveOperation.End)
        cursor.insertBlock()
        key = uid()
        state.features["nodes"][key] = {"kind": "generated", "collection": "endnotes", "label": "All"}
        fmt = QTextCharFormat()
        fmt.setAnchorNames([ANCHOR + key])
        cursor.insertText("Endnotes\n" + "\n".join(f"{index}. {node['text']}" for index, node in enumerate(endnotes, 1)), fmt)
    state.title = (docx_doc.core_properties.title or "").strip() or Path(
        source_name).stem or "Untitled"
    try:
        section = docx_doc.sections[0]
        page = state.page
        if section.page_width is not None:
            width_mm = section.page_width.mm
            height_mm = section.page_height.mm if section.page_height else 297.0
            from .models import PAPER_SIZES_MM
            for name, dims in PAPER_SIZES_MM.items():
                if abs(width_mm - dims[0]) < 2 and abs(height_mm - dims[1]) < 2:
                    page.paper = name
                    page.landscape = False
                    break
                if abs(width_mm - dims[1]) < 2 and abs(height_mm - dims[0]) < 2:
                    page.paper = name
                    page.landscape = True
                    break
            else:
                page.landscape = width_mm > height_mm
        if section.top_margin is not None:
            page.top_mm = section.top_margin.mm
        if section.bottom_margin is not None:
            page.bottom_mm = section.bottom_margin.mm
        if section.left_margin is not None:
            page.left_mm = section.left_margin.mm
        if section.right_margin is not None:
            page.right_mm = section.right_margin.mm
        sect_pr = section._sectPr
        cols = sect_pr.find(qn("w:cols"))
        if cols is not None and cols.get(qn("w:num")):
            page.columns = max(1, min(3, int(cols.get(qn("w:num")))))
        if cols is not None and cols.get(qn("w:space")):
            page.gutter_mm = int(cols.get(qn("w:space"))) / 56.6929
        line_numbers = sect_pr.find(qn("w:lnNumType"))
        if line_numbers is not None:
            restart = line_numbers.get(qn("w:restart"), "continuous")
            page.line_numbering = {
                "newPage": "restart_page",
                "newSection": "restart_section",
            }.get(restart, "continuous")
            try:
                page.line_number_start = max(
                    1, int(line_numbers.get(qn("w:start"), "1")))
                page.line_number_count_by = max(
                    1, int(line_numbers.get(qn("w:countBy"), "1")))
                page.line_number_distance_mm = max(
                    0.0,
                    int(line_numbers.get(qn("w:distance"), "283"))
                    / 56.6929)
            except ValueError:
                page.line_numbering = "none"
        settings_root = docx_doc.settings._element
        auto_hyphen = settings_root.find(qn("w:autoHyphenation"))
        if (auto_hyphen is not None
                and auto_hyphen.get(qn("w:val"), "1") not in (
                    "0", "false", "off")):
            page.hyphenation = "automatic"
            caps = settings_root.find(qn("w:doNotHyphenateCaps"))
            page.hyphenate_caps = not (
                caps is not None
                and caps.get(qn("w:val"), "1") not in (
                    "0", "false", "off"))
        limit = settings_root.find(qn("w:consecutiveHyphenLimit"))
        if limit is not None:
            try:
                page.consecutive_hyphen_limit = max(
                    0, int(limit.get(qn("w:val"), "0")))
            except ValueError:
                page.consecutive_hyphen_limit = 0
        try:
            header_text = _read_hf_text(section.header)
            footer_text = _read_hf_text(section.footer)
            page.header = header_text
            page.footer = footer_text
        except Exception:
            pass
        try:
            page.validate()
        except Exception:
            state.page = PageSettings()
            warnings.append("Page settings were out of range and were reset.")
    except Exception:
        pass
    state.html = document.toHtml()
    return state, warnings


def _import_table(cursor, table, warnings: list[str]) -> None:
    try:
        rows = len(table.rows)
        cols = len(table.columns)
        if rows == 0 or cols == 0:
            return
        fmt = QTextTableFormat()
        fmt.setCellPadding(6.0)
        fmt.setBorder(0.75)
        fmt.setBorderStyle(QTextFrameFormat.BorderStyle.BorderStyle_Solid)
        cursor.insertTable(rows, cols, fmt)
        qtable = cursor.currentTable()
        if not isinstance(qtable, QTextTable):
            return
        seen_cells: set[int] = set()
        merged = False
        for r in range(rows):
            for c in range(cols):
                try:
                    source_cell = table.cell(r, c)
                except Exception:
                    continue
                tc_id = id(source_cell._tc)
                if tc_id in seen_cells:
                    merged = True
                    continue
                seen_cells.add(tc_id)
                try:
                    if source_cell.grid_span > 1:
                        merged = True
                except Exception:
                    pass
                tc = source_cell._tc
                tc_pr = tc.find(qn("w:tcPr"))
                if tc_pr is not None and tc_pr.find(qn("w:vMerge")) is not None:
                    merged = True
                target_cell = qtable.cellAt(r, c)
                cell_cursor = target_cell.firstCursorPosition()
                first_para = True
                for para in source_cell.paragraphs:
                    text = para.text
                    if not first_para:
                        cell_cursor.insertBlock()
                    first_para = False
                    if text:
                        cell_cursor.insertText(text)
        if merged:
            warnings.append(
                "Some merged table cells were simplified to a plain grid.")
        cursor.movePosition(QTextCursor.MoveOperation.End)
    except Exception:
        warnings.append("A table could not be fully imported; content was simplified.")


def read_docx(path: Path) -> tuple[DocumentState, list[str]]:
    path = Path(path)
    try:
        data = path.read_bytes()
    except OSError as exc:
        raise DocxError(f"Could not read the file: {exc}") from exc
    archive = _preflight_zip(data)
    try:
        names = set(archive.namelist())
        if "word/document.xml" not in names:
            raise DocxError("This DOCX has no document body.")
        meta_info = archive.getinfo(META_PART) if META_PART in names else None
        if meta_info is not None and meta_info.file_size > MAX_METADATA_BYTES:
            raise DocxError("The embedded Folio metadata is too large to "
                            "open safely.")
        members = {}
        for info in archive.infolist():
            try:
                members[info.filename] = archive.read(info.filename)
            except (zipfile.BadZipFile, RuntimeError, NotImplementedError,
                    OSError) as exc:
                raise DocxError(
                    f"Could not read '{info.filename}' inside the DOCX: "
                    f"{exc}") from exc
        embedded = _read_embedded_state(archive)
    finally:
        archive.close()
    warnings: list[str] = []
    if embedded is not None:
        try:
            if embedded.get("schema_version") != 1:
                raise ValueError("unsupported schema")
            expected = embedded.get("word_digest", "")
            actual = _word_digest(members)
            if expected != actual:
                raise ValueError("digest mismatch")
            state = DocumentState.from_dict(embedded.get("state"))
            return state, warnings
        except Exception:
            warnings.append(
                "This file was modified outside Folio since Folio last saved "
                "it; the visible Word content was imported instead of the "
                "embedded Folio state."
            )
    state, import_warnings = _import_body(data, path.name)
    warnings.extend(import_warnings)
    return state, warnings
