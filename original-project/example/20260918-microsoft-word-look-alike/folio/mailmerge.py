from __future__ import annotations

import copy
import csv
import re
import uuid
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path

from PySide6.QtCore import QBuffer, QIODevice, QRectF, Qt
from PySide6.QtGui import QImage, QPainter, QTextBlockFormat, QTextCursor, QTextDocumentFragment, QTextFormat, QTextImageFormat

from .editor import SafeDocument
from .models import DocumentState, PageSettings

FIELD_RE = re.compile(r"\{\{([^{}]+)\}\}")


class MergeError(RuntimeError):
    pass


class MergeSkipped(MergeError):
    pass


OPERATORS = ("equals", "not equals", "contains", "starts with", "is empty", "is not empty", "greater than", "less than")
ADDRESS_FIELDS = ("FullName", "Title", "FirstName", "LastName", "Company", "Address1", "Address2", "City", "State", "PostalCode", "Country")


@dataclass(frozen=True)
class Condition:
    field: str
    operator: str = "equals"
    value: str = ""

    def evaluate(self, record: dict[str, str]) -> bool:
        if self.operator not in OPERATORS:
            raise MergeError(f"Unknown comparison: {self.operator}")
        if self.field not in record:
            raise MergeError(f"Missing field: {self.field}")
        left, right = str(record[self.field]).strip(), str(self.value).strip()
        if self.operator == "is empty":
            return not left
        if self.operator == "is not empty":
            return bool(left)
        if self.operator in ("greater than", "less than"):
            try:
                a, b = Decimal(left), Decimal(right)
                if not a.is_finite() or not b.is_finite():
                    raise InvalidOperation
            except InvalidOperation as exc:
                raise MergeError(f"Numeric comparison requires finite numbers in {self.field}.") from exc
            return a > b if self.operator == "greater than" else a < b
        left, right = left.casefold(), right.casefold()
        return {"equals": lambda: left == right,
                "not equals": lambda: left != right,
                "contains": lambda: right in left,
                "starts with": lambda: left.startswith(right)}[self.operator]()


@dataclass
class RecipientSet:
    records: list[dict[str, str]]
    excluded: set[int] = field(default_factory=set)
    filters: list[Condition] = field(default_factory=list)
    sort_field: str = ""
    descending: bool = False

    def active(self) -> list[tuple[int, dict[str, str]]]:
        rows = [(index, record) for index, record in enumerate(self.records)
                if index not in self.excluded and all(rule.evaluate(record) for rule in self.filters)]
        if self.sort_field:
            if any(self.sort_field not in record for _, record in rows):
                raise MergeError(f"Missing sort field: {self.sort_field}")
            def key(row):
                value = row[1][self.sort_field].strip()
                try:
                    number = Decimal(value)
                    if number.is_finite():
                        return (0, number)
                except InvalidOperation:
                    pass
                return (1, value.casefold())
            rows.sort(key=key, reverse=self.descending)
        return rows


def mapped_values(record: dict[str, str], mapping: dict[str, str] | None = None) -> dict[str, str]:
    result = {name: str(value) for name, value in record.items()}
    for target, source in (mapping or {}).items():
        if source:
            if source not in record:
                raise MergeError(f"Mapped column is missing: {source}")
            result[target] = str(record[source])
    return result


def address_block(values: dict[str, str]) -> str:
    name = values.get("FullName", "").strip() or " ".join(
        values.get(key, "").strip() for key in ("Title", "FirstName", "LastName") if values.get(key, "").strip())
    city = ", ".join(values.get(key, "").strip() for key in ("City", "State") if values.get(key, "").strip())
    locality = " ".join(part for part in (city, values.get("PostalCode", "").strip()) if part)
    return "\n".join(part for part in (name, values.get("Company", "").strip(),
        values.get("Address1", "").strip(), values.get("Address2", "").strip(),
        locality, values.get("Country", "").strip()) if part)


def greeting_line(values: dict[str, str]) -> str:
    name = values.get("FullName", "").strip() or " ".join(
        values.get(key, "").strip() for key in ("Title", "LastName") if values.get(key, "").strip())
    name = name or values.get("FirstName", "").strip()
    return f"Dear {name}," if name else "Dear Sir or Madam,"


def rule_token(kind: str, condition: Condition, then_text: str = "", else_text: str = "") -> str:
    if kind not in ("IF", "SKIPIF") or condition.operator not in OPERATORS or not condition.field:
        raise MergeError("Choose a rule, field, and valid comparison.")
    parts = [kind, condition.field, condition.operator, condition.value]
    if kind == "IF":
        parts += [then_text, else_text]
    if any(any(char in part for char in "|{}\r\n") for part in parts):
        raise MergeError("Rule values cannot contain pipes, braces, or newlines.")
    return "{{" + "|".join(parts) + "}}"


def _token_value(token: str, values: dict[str, str]) -> str:
    if token.startswith(("IF|", "SKIPIF|")):
        parts = token.split("|")
        if len(parts) != (6 if parts[0] == "IF" else 4):
            raise MergeError("Invalid merge rule. Insert it using Rules.")
        matches = Condition(parts[1], parts[2], parts[3]).evaluate(values)
        if parts[0] == "SKIPIF":
            if matches:
                raise MergeSkipped("Record excluded by a Skip If rule.")
            return ""
        return parts[4] if matches else parts[5]
    if token == "AddressBlock":
        return address_block(values)
    if token == "GreetingLine":
        return greeting_line(values)
    if token not in values:
        raise MergeError(f"Missing field: {token}")
    return values[token]


def _read_csv(path: Path) -> list[dict[str, str]]:
    records: list[dict[str, str]] = []
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = reader.fieldnames or []
        _validate_headers(headers)
        for row in reader:
            record = {str(k).strip(): ("" if v is None else str(v))
                      for k, v in row.items() if k is not None}
            if any(value.strip() for value in record.values()):
                records.append(record)
    return records


def _read_xlsx(path: Path) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise MergeError("openpyxl is required for XLSX merge sources.") from exc
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise MergeError(f"Could not read the XLSX file: {exc}") from exc
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            raise MergeError("The XLSX sheet is empty.")
        headers = ["" if cell is None else str(cell).strip()
                   for cell in header_row]
        _validate_headers(headers)
        records: list[dict[str, str]] = []
        for row in rows:
            record = {}
            for index, header in enumerate(headers):
                value = row[index] if index < len(row) else None
                record[header] = "" if value is None else str(value)
            if any(value.strip() for value in record.values()):
                records.append(record)
        return records
    finally:
        workbook.close()


def _validate_headers(headers: list) -> None:
    cleaned = [str(h).strip() for h in headers]
    if not cleaned or any(not h for h in cleaned):
        raise MergeError("The merge source has an empty header column.")
    if len(set(cleaned)) != len(cleaned):
        raise MergeError("The merge source has duplicate header names.")


def read_records(path: Path) -> list[dict[str, str]]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv(path)
    if suffix == ".xlsx":
        return _read_xlsx(path)
    raise MergeError("Merge sources must be .csv or .xlsx files.")


def discover_fields(document) -> list[str]:
    text = document.toPlainText()
    seen: list[str] = []
    for match in FIELD_RE.finditer(text):
        name = match.group(1).strip()
        if name.startswith(("IF|", "SKIPIF|")):
            name = name.split("|")[1]
        if name in ("AddressBlock", "GreetingLine"):
            continue
        if name and name not in seen:
            seen.append(name)
    return seen


def merge_document(document, values: dict[str, str],
                   mapping: dict[str, str] | None = None) -> SafeDocument:
    clone = SafeDocument()
    clone.setDefaultFont(document.defaultFont())
    clone.setHtml(document.toHtml())
    text = clone.toPlainText()
    matches = list(FIELD_RE.finditer(text))
    values = mapped_values(values, mapping)
    replacements = [_token_value(match.group(1).strip(), values) for match in matches]
    for match, replacement in reversed(list(zip(matches, replacements))):
        start = len(text[:match.start()].encode("utf-16-le")) // 2
        end = len(text[:match.end()].encode("utf-16-le")) // 2
        cursor = QTextCursor(clone)
        cursor.setPosition(start)
        cursor.setPosition(end, QTextCursor.MoveMode.KeepAnchor)
        cursor.insertText(replacement, cursor.charFormat())
    return clone


def _sanitize_title(title: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9 _-]+", "", title).strip()
    cleaned = cleaned.replace(" ", "_")
    return cleaned or "merged"


def merge_to_directory(document, state: DocumentState,
                       records: list[dict[str, str]],
                       directory: Path,
                       progress=None, cancelled=None) -> list[Path]:
    from .docx_io import write_docx

    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    fields = discover_fields(document)
    missing = [f for f in fields
               if any(f not in record for record in records)]
    if missing:
        raise MergeError(
            "The merge source is missing required fields: "
            + ", ".join(missing)
        )
    outputs: list[Path] = []
    base = _sanitize_title(state.title)
    for index, record in enumerate(records):
        if cancelled is not None and cancelled():
            break
        try:
            merged = merge_document(document, record)
        except MergeSkipped:
            if progress is not None:
                progress(index + 1, len(records))
            continue
        new_state = copy.deepcopy(state)
        new_state.id = str(__import__("uuid").uuid4())
        new_state.comments = []
        new_state.revisions = []
        new_state.toc = None
        new_state.track_changes = False
        if fields:
            first = fields[0]
            if record.get(first):
                new_state.title = f"{state.title} - {record[first]}"
        candidate = directory / f"{index + 1:03d}_{base}.docx"
        suffix = 1
        while candidate.exists():
            candidate = directory / f"{index + 1:03d}_{base}_{suffix}.docx"
            suffix += 1
        write_docx(candidate, new_state, merged)
        outputs.append(candidate)
        if progress is not None:
            progress(index + 1, len(records))
    return outputs


@dataclass(frozen=True)
class LabelTemplate:
    name: str
    paper: str
    columns: int
    rows: int
    width_mm: float
    height_mm: float
    left_mm: float
    top_mm: float
    gap_x_mm: float = 0.0
    gap_y_mm: float = 0.0

    @property
    def capacity(self) -> int:
        return self.columns * self.rows

    def settings(self) -> PageSettings:
        return PageSettings(paper=self.paper, top_mm=0, bottom_mm=0,
                            left_mm=0, right_mm=0, header="", footer="")

    def rectangles(self) -> list[QRectF]:
        width, height = self.settings().size_mm()
        rectangles = [QRectF(self.left_mm + column * (self.width_mm + self.gap_x_mm),
                             self.top_mm + row * (self.height_mm + self.gap_y_mm),
                             self.width_mm, self.height_mm)
                      for row in range(self.rows) for column in range(self.columns)]
        if (self.rows < 1 or self.columns < 1 or self.width_mm <= 4 or self.height_mm <= 4
                or any(rect.left() < 0 or rect.top() < 0 or rect.right() > width + .01
                       or rect.bottom() > height + .01 for rect in rectangles)):
            raise MergeError("Label dimensions do not fit on the sheet.")
        return rectangles


LABEL_TEMPLATES = {
    "Avery 5160 / 8160 (30 per Letter sheet)": LabelTemplate(
        "Avery 5160 / 8160 (30 per Letter sheet)", "Letter", 3, 10,
        66.675, 25.4, 4.7625, 12.7, 3.175),
    "Avery L7160 (21 per A4 sheet)": LabelTemplate(
        "Avery L7160 (21 per A4 sheet)", "A4", 3, 7,
        63.5, 38.1, 7.25, 15.15, 2.5),
    "Avery L7163 (14 per A4 sheet)": LabelTemplate(
        "Avery L7163 (14 per A4 sheet)", "A4", 2, 7,
        99.1, 38.1, 4.65, 15.15, 2.5),
}
ENVELOPE_TEMPLATES = ("Envelope #10", "Envelope DL", "Envelope C5")


def envelope_settings(name: str) -> PageSettings:
    if name not in ENVELOPE_TEMPLATES:
        raise MergeError("Unknown envelope format.")
    page = PageSettings(paper=name, landscape=True, top_mm=40, left_mm=80,
                        right_mm=15, bottom_mm=15, header="", footer="")
    page.validate()
    return page


def checked_mailing(value) -> dict:
    if value is None:
        return {}
    if not isinstance(value, dict):
        raise MergeError("Mailing settings must be an object.")
    mapping = value.get("mapping", {})
    if (not isinstance(mapping, dict) or len(mapping) > 500
            or any(not isinstance(key, str) or not isinstance(source, str)
                   or len(key) > 200 or len(source) > 200 for key, source in mapping.items())):
        raise MergeError("Invalid mailing field mapping.")
    label = value.get("label_template", "")
    if not isinstance(label, str) or (label and label not in LABEL_TEMPLATES):
        raise MergeError("Unknown label template.")
    return {"mapping": dict(mapping), "label_template": label}


def prepared_records(document, records, mapping=None):
    prepared = []
    for index, record in enumerate(records):
        try:
            prepared.append((index, record, merge_document(document, record, mapping)))
        except MergeSkipped:
            continue
        except MergeError as exc:
            raise MergeError(f"Record {index + 1}: {exc}") from exc
    if not prepared:
        raise MergeError("No recipients remain after exclusions and Skip If rules.")
    return prepared


def compound_document(documents) -> SafeDocument:
    result = SafeDocument()
    cursor = QTextCursor(result)
    for index, document in enumerate(documents):
        if index:
            cursor.movePosition(QTextCursor.MoveOperation.End)
            fmt = QTextBlockFormat()
            fmt.setPageBreakPolicy(QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore)
            cursor.insertBlock(fmt)
        else:
            result.setDefaultFont(document.defaultFont())
        start = cursor.position()
        cursor.insertFragment(QTextDocumentFragment(document))
        if index:
            marker = QTextCursor(result)
            marker.setPosition(start)
            fmt = marker.blockFormat()
            fmt.setPageBreakPolicy(QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore)
            marker.setBlockFormat(fmt)
    return result


def label_sheets(documents, template: LabelTemplate) -> SafeDocument:
    from .publishing import mm_to_px

    rectangles = template.rectangles()
    width_mm, height_mm = template.settings().size_mm()
    width, height = mm_to_px(width_mm), mm_to_px(height_mm)
    sheets = []
    for offset in range(0, len(documents), template.capacity):
        image = QImage(round(width * 2), round(height * 2), QImage.Format.Format_ARGB32)
        image.fill(Qt.GlobalColor.white)
        painter = QPainter(image)
        painter.scale(2, 2)
        try:
            for rect_mm, source in zip(rectangles, documents[offset:offset + template.capacity]):
                rect = QRectF(mm_to_px(rect_mm.x() + 2), mm_to_px(rect_mm.y() + 2),
                              mm_to_px(rect_mm.width() - 4), mm_to_px(rect_mm.height() - 4))
                doc = SafeDocument()
                doc.setDefaultFont(source.defaultFont())
                doc.setHtml(source.toHtml())
                doc.setDocumentMargin(0)
                doc.setTextWidth(rect.width())
                if doc.size().height() > rect.height() + .5 or doc.idealWidth() > rect.width() + .5:
                    raise MergeError("A recipient's content does not fit on its label. Shorten the text or reduce its font size.")
                painter.save()
                painter.translate(rect.topLeft())
                doc.drawContents(painter, QRectF(0, 0, rect.width(), rect.height()))
                painter.restore()
        finally:
            painter.end()
        buffer = QBuffer()
        buffer.open(QIODevice.OpenModeFlag.WriteOnly)
        image.save(buffer, "PNG")
        uri = "data:image/png;base64," + bytes(buffer.data().toBase64()).decode("ascii")
        sheet = SafeDocument()
        sheet.setDocumentMargin(0)
        cursor = QTextCursor(sheet)
        fmt = QTextBlockFormat()
        fmt.setTopMargin(0)
        fmt.setBottomMargin(0)
        fmt.setLineHeight(height - .05, QTextBlockFormat.LineHeightTypes.FixedHeight.value)
        cursor.setBlockFormat(fmt)
        image_format = QTextImageFormat()
        image_format.setName(uri)
        image_format.setWidth(width - .05)
        image_format.setHeight(height - .05)
        cursor.insertImage(image_format)
        sheets.append(sheet)
    return compound_document(sheets)


def merged_state(state, title=None) -> DocumentState:
    result = copy.deepcopy(state)
    result.id = str(uuid.uuid4())
    result.title = title or state.title
    result.comments = []
    result.revisions = []
    result.toc = None
    result.track_changes = False
    result.features.pop("mailing", None)
    return result


def _unique_output(directory: Path, stem: str, extension: str) -> Path:
    candidate = directory / f"{stem}.{extension}"
    index = 1
    while candidate.exists():
        candidate = directory / f"{stem}_{index}.{extension}"
        index += 1
    return candidate


def _check_output_features(state, output, labels=""):
    features = state.features
    if ((features.get("nodes") or features.get("objects") or features.get("sections"))
            and (output != "docx" or labels)):
        raise MergeError("This template contains notes, linked objects, or section settings. "
                         "Export individual DOCX files without label sheets to preserve these features.")


def merge_batch(document, state, records, directory, output="docx", mapping=None,
                labels="", email_field="Email", subject="", progress=None, cancelled=None):
    from .docx_io import write_docx
    from .publishing import write_pdf

    if output not in ("docx", "pdf", "compound_docx", "compound_pdf", "eml"):
        raise MergeError("Choose DOCX, PDF, a compound document, or email payloads.")
    _check_output_features(state, output, labels)
    prepared = prepared_records(document, records, mapping)
    if labels and labels not in LABEL_TEMPLATES:
        raise MergeError("Unknown label template.")
    if output == "eml":
        if labels:
            raise MergeError("Email payloads require a letter template, not label sheets.")
        for index, record, _ in prepared:
            recipient = mapped_values(record, mapping).get(email_field, "").strip()
            if not re.fullmatch(r"[^\s@<>;,]+@[^\s@<>;,]+\.[^\s@<>;,]+", recipient):
                raise MergeError(f"Record {index + 1}: invalid email address in {email_field}.")
        if any(char in subject for char in "\r\n\x00"):
            raise MergeError("Email subject cannot contain newlines.")
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    base = _sanitize_title(state.title)
    batch_state = merged_state(state)
    if labels:
        template = LABEL_TEMPLATES[labels]
        batch_state.page = template.settings()
        batch_state.features = {"nodes": {}, "objects": [], "sections": []}
        batches = [(index, {}, label_sheets([row[2] for row in prepared[index:index + template.capacity]], template))
                   for index in range(0, len(prepared), template.capacity)]
    else:
        batches = prepared
    if output.startswith("compound_"):
        combined = compound_document([row[2] for row in batches])
        batches = [(0, {}, combined)]
        batch_state.features = {"nodes": {}, "objects": [], "sections": []}
    outputs = []
    extension = output.removeprefix("compound_")
    for ordinal, (source_index, record, doc) in enumerate(batches):
        if cancelled is not None and cancelled():
            break
        candidate = _unique_output(directory, f"{ordinal + 1:03d}_{base}", extension)
        new_state = merged_state(batch_state)
        doc.folio_state = new_state
        if extension == "docx":
            write_docx(candidate, new_state, doc)
        elif extension == "pdf":
            write_pdf(candidate, doc, new_state.page, new_state.title)
        else:
            message = EmailMessage(policy=SMTP)
            message["To"] = mapped_values(record, mapping)[email_field]
            message["Subject"] = subject or state.title
            message["X-Unsent"] = "1"
            message.set_content(doc.toPlainText())
            message.add_alternative(doc.toHtml(), subtype="html")
            with candidate.open("xb") as handle:
                handle.write(message.as_bytes())
        outputs.append(candidate)
        if progress is not None:
            progress(ordinal + 1, len(batches))
    return outputs


def print_batch(document, state, records, printer, mapping=None, labels="") -> int:
    from .publishing import Publication

    _check_output_features(state, "print", labels)
    documents = [row[2] for row in prepared_records(document, records, mapping)]
    settings = state.page
    if labels:
        if labels not in LABEL_TEMPLATES:
            raise MergeError("Unknown label template.")
        template = LABEL_TEMPLATES[labels]
        combined = label_sheets(documents, template)
        settings = template.settings()
    else:
        combined = compound_document(documents)
    return Publication(combined, settings, state.title).paint(printer)
