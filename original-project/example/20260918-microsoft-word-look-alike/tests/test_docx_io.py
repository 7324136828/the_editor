import json
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from PySide6.QtGui import (QColor, QImage, QTextCharFormat, QTextCursor, QTextFormat,
                           QTextImageFormat)

import docx
from docx.oxml.ns import qn

from folio.docx_io import (DocxError, META_PART, _word_digest, atomic_write,
                           read_docx, write_docx)
from folio.editor import (COLUMN_BREAK_MARKER, COLUMN_BREAK_URL, SafeDocument,
                          _image_to_data_uri)
from folio.models import Comment, DocumentState, PageSettings, Revision


@pytest.fixture
def document(qapp):
    return SafeDocument()


def _make_state(**kwargs):
    state = DocumentState(title="Round Trip")
    for key, value in kwargs.items():
        setattr(state, key, value)
    return state


def test_round_trip_heading_and_body(document, tmp_path):
    document.setHtml(
        "<h1>Big Heading</h1><p>Hello <b>bold</b> world</p>")
    state = _make_state()
    path = tmp_path / "doc.docx"
    write_docx(path, state, document)
    loaded, warnings = read_docx(path)
    plain = loaded.html
    assert "Big Heading" in plain
    assert "bold" in plain
    docx_doc = docx.Document(str(path))
    texts = [p.text for p in docx_doc.paragraphs]
    assert "Big Heading" in texts
    assert any("bold" in t for t in texts)
    styles = {p.style.name for p in docx_doc.paragraphs}
    assert "Heading 1" in styles


def test_round_trip_table(document, tmp_path):
    document.setHtml(
        "<table border='1'><tr><td>A1</td><td>B1</td></tr>"
        "<tr><td>A2</td><td>B2</td></tr></table>")
    path = tmp_path / "table.docx"
    write_docx(path, _make_state(), document)
    docx_doc = docx.Document(str(path))
    assert len(docx_doc.tables) == 1
    table = docx_doc.tables[0]
    assert table.cell(0, 0).text == "A1"
    assert table.cell(1, 1).text == "B2"


def test_merged_table_cell_exports_once(document, tmp_path):
    cursor = QTextCursor(document)
    table = cursor.insertTable(2, 2)
    cell = table.cellAt(0, 0)
    c1 = cell.firstCursorPosition()
    c1.insertText("merged title")
    c1.insertBlock()
    c1.insertText("second line")
    table.mergeCells(0, 0, 1, 2)
    path = tmp_path / "merged.docx"
    write_docx(path, _make_state(), document)
    docx_doc = docx.Document(str(path))
    dt = docx_doc.tables[0]
    xml = dt._tbl.xml
    assert xml.count("merged title") == 1
    assert xml.count("second line") == 1
    cell = dt.cell(0, 0)
    assert len(cell.paragraphs) == 2
    assert cell.paragraphs[0].text == "merged title"
    assert cell.paragraphs[1].text == "second line"
    assert dt.cell(0, 1)._tc is cell._tc


def test_embedded_image_round_trip(document, tmp_path):
    cursor = QTextCursor(document)
    image = QImage(40, 30, QImage.Format.Format_RGB32)
    image.fill(QColor("#224466"))
    fmt = QTextImageFormat()
    fmt.setName(_image_to_data_uri(image))
    fmt.setWidth(40)
    fmt.setHeight(30)
    cursor.insertImage(fmt)
    path = tmp_path / "img.docx"
    write_docx(path, _make_state(), document)
    with zipfile.ZipFile(path) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        assert media
    loaded, _ = read_docx(path)
    assert "￼" in loaded.html or "img" in loaded.html.lower() or "data:image" in loaded.html


def test_page_settings_round_trip(document, tmp_path):
    document.setPlainText("page")
    state = _make_state()
    state.page = PageSettings(paper="Letter", landscape=True, columns=2,
                              top_mm=20.0, header="Doc {title}",
                              footer="Page {page} of {pages}")
    path = tmp_path / "page.docx"
    write_docx(path, state, document)
    docx_doc = docx.Document(str(path))
    section = docx_doc.sections[0]
    assert section.page_width.mm == pytest.approx(279.4, abs=1.5)
    assert section.page_height.mm == pytest.approx(215.9, abs=1.5)
    assert section.top_margin.mm == pytest.approx(20.0, abs=0.5)
    cols = section._sectPr.find(qn("w:cols"))
    assert cols is not None and cols.get(qn("w:num")) == "2"
    footer_xml = section.footer._element.xml
    assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml
    header_xml = section.header._element.xml
    assert "Doc" in header_xml


def test_line_numbers_and_hyphenation_exported(document, tmp_path):
    document.setPlainText("layout")
    state = _make_state()
    state.page.line_numbering = "restart_page"
    state.page.line_number_start = 3
    state.page.line_number_count_by = 2
    state.page.line_number_distance_mm = 6.0
    state.page.hyphenation = "automatic"
    state.page.hyphenate_caps = True
    state.page.consecutive_hyphen_limit = 2
    path = tmp_path / "layout.docx"
    write_docx(path, state, document)
    docx_doc = docx.Document(str(path))
    line_numbers = docx_doc.sections[0]._sectPr.find(qn("w:lnNumType"))
    assert line_numbers is not None
    assert line_numbers.get(qn("w:restart")) == "newPage"
    assert line_numbers.get(qn("w:start")) == "3"
    assert line_numbers.get(qn("w:countBy")) == "2"
    settings = docx_doc.settings._element
    assert settings.find(qn("w:autoHyphenation")) is not None
    assert settings.find(qn("w:doNotHyphenateCaps")) is None
    limit = settings.find(qn("w:consecutiveHyphenLimit"))
    assert limit is not None and limit.get(qn("w:val")) == "2"
    loaded, _warnings = read_docx(path)
    assert loaded.page.line_numbering == "restart_page"
    assert loaded.page.hyphenation == "automatic"


def test_column_break_exports_as_column_break(document, tmp_path):
    document.setPlainText("before")
    cursor = QTextCursor(document)
    cursor.movePosition(QTextCursor.MoveOperation.End)
    cursor.insertBlock()
    block_fmt = cursor.blockFormat()
    block_fmt.setPageBreakPolicy(
        QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore)
    cursor.setBlockFormat(block_fmt)
    marker = QTextCharFormat()
    marker.setAnchor(True)
    marker.setAnchorHref(COLUMN_BREAK_URL)
    cursor.insertText(COLUMN_BREAK_MARKER, marker)
    cursor.insertText("after")
    path = tmp_path / "column-break.docx"
    write_docx(path, _make_state(), document)
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert 'w:type="column"' in xml
    assert COLUMN_BREAK_MARKER not in xml


def test_comments_revisions_metadata_round_trip(document, tmp_path):
    document.setPlainText("0123456789")
    state = _make_state()
    state.comments.append(Comment(start=0, end=4, quote="0123",
                                  text="note", author="A"))
    state.revisions.append(Revision(start=2, length=3, before_text="234",
                                    after_text="X", before_html="234",
                                    after_html="X", author="B"))
    path = tmp_path / "meta.docx"
    write_docx(path, state, document)
    loaded, _ = read_docx(path)
    assert loaded.comments and loaded.comments[0].text == "note"
    assert loaded.revisions and loaded.revisions[0].after_text == "X"


def test_list_numbering_exported(document, tmp_path):
    document.setHtml("<ol><li>first</li><li>second</li></ol>")
    path = tmp_path / "list.docx"
    write_docx(path, _make_state(), document)
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        assert "word/numbering.xml" in names
        rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
        assert rels.count("relationships/numbering") == 1
    docx_doc = docx.Document(str(path))
    assert docx_doc.part.numbering_part is not None


def test_page_break_exported(document, tmp_path):
    document.setHtml("<p>one</p><p>two</p>")
    cursor = QTextCursor(document)
    cursor.setPosition(4)
    cursor.insertBlock()
    from PySide6.QtGui import QTextFormat
    fmt = cursor.blockFormat()
    fmt.setPageBreakPolicy(QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore)
    cursor.setBlockFormat(fmt)
    path = tmp_path / "brk.docx"
    write_docx(path, _make_state(), document)
    docx_doc = docx.Document(str(path))
    breaks = [
        p.paragraph_format.page_break_before for p in docx_doc.paragraphs]
    assert any(breaks)


def test_write_docx_does_not_mutate_state(document, tmp_path):
    document.setPlainText("abc")
    state = _make_state()
    state.html = ""
    state.track_changes = True
    write_docx(tmp_path / "x.docx", state, document)
    assert state.html == ""
    assert state.track_changes is True


def test_state_snapshot_matches_document(document, tmp_path):
    document.setPlainText("live body")
    state = _make_state()
    state.html = "STALE"
    path = tmp_path / "snap.docx"
    write_docx(path, state, document)
    with zipfile.ZipFile(path) as z:
        meta = json.loads(z.read(META_PART).decode("utf-8"))
    assert "live body" in meta["state"]["html"]
    assert meta["state"]["track_changes"] == state.track_changes


def test_track_changes_state_round_trip(document, tmp_path):
    document.setPlainText("x")
    state = _make_state()
    state.track_changes = True
    path = tmp_path / "track.docx"
    write_docx(path, state, document)
    loaded, _ = read_docx(path)
    assert loaded.track_changes is True
    assert state.track_changes is True


def test_oversized_metadata_write_preserves_existing(document, tmp_path,
                                                     monkeypatch):
    document.setPlainText("x")
    state = _make_state()
    state.html = "y" * 5000
    path = tmp_path / "keep.docx"
    path.write_bytes(b"ORIGINAL BYTES")
    import folio.docx_io as dio
    monkeypatch.setattr(dio, "MAX_METADATA_BYTES", 64)
    with pytest.raises(DocxError):
        write_docx(path, state, document)
    assert path.read_bytes() == b"ORIGINAL BYTES"


def test_write_docx_rejects_invalid_page(document, tmp_path):
    document.setPlainText("x")
    state = _make_state()
    state.page.columns = 5
    with pytest.raises(DocxError):
        write_docx(tmp_path / "bad.docx", state, document)


def test_corrupt_zip_member_raises_docx_error(document, tmp_path):
    document.setPlainText("x")
    path = tmp_path / "corrupt.docx"
    write_docx(path, _make_state(), document)
    data = bytearray(path.read_bytes())
    middle = len(data) // 2
    data[middle] ^= 0xFF
    data[middle + 1] ^= 0xFF
    path.write_bytes(bytes(data))
    with pytest.raises(DocxError):
        read_docx(path)


def test_external_inline_page_break_single_page(tmp_path):
    from docx.enum.text import WD_BREAK
    from folio.publishing import Publication
    docx_doc = docx.Document()
    p = docx_doc.add_paragraph()
    p.add_run("A")
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)
    run.add_text("B")
    docx_doc.add_paragraph("C")
    path = tmp_path / "inline_break.docx"
    docx_doc.save(str(path))
    loaded, _ = read_docx(path)
    doc = SafeDocument()
    doc.setHtml(loaded.html)
    breaks = 0
    block = doc.begin()
    while block.isValid():
        if block.blockFormat().pageBreakPolicy() == (
                QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore):
            breaks += 1
        block = block.next()
    assert breaks == 1
    publication = Publication(doc, loaded.page, "t")
    plain = doc.toPlainText()
    b_page = publication.page_for_position(plain.find("B"))
    c_page = publication.page_for_position(plain.find("C"))
    assert b_page == c_page == 2


def test_external_modification_falls_back_to_body(document, tmp_path):
    document.setPlainText("original text")
    path = tmp_path / "ext.docx"
    write_docx(path, _make_state(), document)
    with zipfile.ZipFile(path) as z:
        members = {i.filename: z.read(i.filename) for i in z.infolist()}
    body = members["word/document.xml"].decode("utf-8")
    body = body.replace("original text", "EDITED EXTERNALLY")
    members["word/document.xml"] = body.encode("utf-8")
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in members.items():
            zout.writestr(name, data)
    path.write_bytes(buffer.getvalue())
    loaded, warnings = read_docx(path)
    assert any("modified outside" in w for w in warnings)
    assert "EDITED EXTERNALLY" in loaded.html


def test_invalid_zip_rejected(tmp_path):
    path = tmp_path / "bad.docx"
    path.write_bytes(b"this is not a zip file at all")
    with pytest.raises(DocxError):
        read_docx(path)


def test_oversized_metadata_rejected(document, tmp_path):
    document.setPlainText("x")
    path = tmp_path / "big.docx"
    write_docx(path, _make_state(), document)
    with zipfile.ZipFile(path) as z:
        members = {i.filename: z.read(i.filename) for i in z.infolist()}
    members[META_PART] = b"x" * (17 * 1024 * 1024)
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_STORED) as zout:
        for name, data in members.items():
            zout.writestr(name, data)
    path.write_bytes(buffer.getvalue())
    with pytest.raises(DocxError):
        read_docx(path)


def test_atomic_write_failure_preserves_old(tmp_path):
    path = tmp_path / "keep.bin"
    path.write_bytes(b"original")
    import folio.docx_io as dio
    original_replace = dio.os.replace
    def failing_replace(src, dst):
        raise OSError("simulated failure")
    dio.os.replace = failing_replace
    try:
        with pytest.raises(OSError):
            atomic_write(path, b"new data")
    finally:
        dio.os.replace = original_replace
    assert path.read_bytes() == b"original"
    leftovers = list(tmp_path.glob("*.tmp"))
    assert not leftovers


def test_external_docx_import_basic(tmp_path):
    docx_doc = docx.Document()
    docx_doc.add_heading("External Title", level=1)
    docx_doc.add_paragraph("plain body text")
    path = tmp_path / "external.docx"
    docx_doc.save(str(path))
    loaded, warnings = read_docx(path)
    assert "External Title" in loaded.html
    assert "plain body text" in loaded.html
    assert warnings


def test_external_empty_footer(tmp_path):
    docx_doc = docx.Document()
    docx_doc.add_paragraph("body")
    path = tmp_path / "nofoot.docx"
    docx_doc.save(str(path))
    loaded, _ = read_docx(path)
    assert loaded.page.footer == ""


def test_external_run_order_mixed_content(tmp_path):
    docx_doc = docx.Document()
    p = docx_doc.add_paragraph()
    r = p.add_run("before")
    r.add_tab()
    r.add_text("after")
    path = tmp_path / "runs.docx"
    docx_doc.save(str(path))
    loaded, _ = read_docx(path)
    html = loaded.html
    assert html.find("before") < html.find("after")


def test_digest_changes_on_body_edit(document, tmp_path):
    document.setPlainText("hash me")
    path = tmp_path / "d.docx"
    write_docx(path, _make_state(), document)
    with zipfile.ZipFile(path) as z:
        members = {i.filename: z.read(i.filename) for i in z.infolist()}
    d1 = _word_digest(members)
    members["word/document.xml"] += b" "
    d2 = _word_digest(members)
    assert d1 != d2
