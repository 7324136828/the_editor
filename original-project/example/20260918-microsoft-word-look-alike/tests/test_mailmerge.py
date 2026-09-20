import csv
from pathlib import Path

import pytest

from folio.docx_io import read_docx
from folio.editor import SafeDocument
from folio.mailmerge import (MergeError, discover_fields, merge_document,
                             merge_to_directory, read_records)
from folio.models import DocumentState


@pytest.fixture
def document(qapp):
    doc = SafeDocument()
    doc.setHtml("<p>Dear {{Name}},</p><p>Your city: {{City}}</p>")
    return doc


def _csv(path: Path, rows: list[dict], bom: bool = True):
    with open(path, "w", encoding="utf-8-sig" if bom else "utf-8",
              newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def test_csv_with_bom(tmp_path):
    path = tmp_path / "data.csv"
    _csv(path, [{"Name": "Ada", "City": "Paris"}])
    records = read_records(path)
    assert records == [{"Name": "Ada", "City": "Paris"}]


def test_xlsx_source(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "City"])
    ws.append(["Grace", "Rome"])
    path = tmp_path / "data.xlsx"
    wb.save(path)
    records = read_records(path)
    assert records == [{"Name": "Grace", "City": "Rome"}]


def test_bad_headers_rejected(tmp_path):
    path = tmp_path / "bad.csv"
    path.write_text("Name,\nAda,x\n", encoding="utf-8")
    with pytest.raises(MergeError):
        read_records(path)


def test_duplicate_headers_rejected(tmp_path):
    path = tmp_path / "dup.csv"
    path.write_text("Name,Name\nAda,Bea\n", encoding="utf-8")
    with pytest.raises(MergeError):
        read_records(path)


def test_unsupported_extension(tmp_path):
    path = tmp_path / "data.txt"
    path.write_text("x")
    with pytest.raises(MergeError):
        read_records(path)


def test_discover_fields(document):
    assert discover_fields(document) == ["Name", "City"]


def test_merge_document_replaces(document):
    merged = merge_document(document, {"Name": "Ada", "City": "Paris"})
    text = merged.toPlainText()
    assert "Ada" in text and "Paris" in text
    assert "{{" not in text


def test_missing_field_raises(document):
    with pytest.raises(MergeError):
        merge_document(document, {"Name": "Ada"})


def test_self_token_value_no_hang(document):
    merged = merge_document(document, {"Name": "{{Name}}", "City": "X"})
    text = merged.toPlainText()
    assert "{{Name}}" in text


def test_other_token_value_not_recursed(document):
    merged = merge_document(document, {"Name": "{{City}}", "City": "Paris"})
    text = merged.toPlainText()
    assert "{{City}}" in text
    assert "Dear Paris" not in text


def test_formatted_split_field(qapp):
    doc = SafeDocument()
    doc.setHtml("<p>Hello {{First<b>Name</b>}}!</p>")
    merged = merge_document(doc, {"FirstName": "Ada"})
    assert "Hello Ada!" == merged.toPlainText().strip()


def test_literal_markup_stays_text(document):
    merged = merge_document(document,
                            {"Name": "<b>literal</b>", "City": "Y"})
    text = merged.toPlainText()
    assert "<b>literal</b>" in text


def test_astral_prefix_offsets(qapp):
    doc = SafeDocument()
    doc.setHtml("<p>\U0001D11E {{Name}} ok</p>")
    merged = merge_document(doc, {"Name": "Zed"})
    text = merged.toPlainText()
    assert "\U0001D11E Zed ok" in text


def test_merge_to_directory(document, tmp_path):
    records = [{"Name": "Ada", "City": "Paris"},
               {"Name": "Bea", "City": "Rome"}]
    state = DocumentState(title="Letter")
    outputs = merge_to_directory(document, state, records, tmp_path / "out")
    assert len(outputs) == 2
    assert outputs[0].name != outputs[1].name
    loaded, _ = read_docx(outputs[0])
    assert "Ada" in loaded.html and "Paris" in loaded.html
    assert "{{Name}}" not in loaded.html
    loaded2, _ = read_docx(outputs[1])
    assert "Bea" in loaded2.html
    assert loaded.id != loaded2.id


def test_merge_existing_filenames_suffixed(document, tmp_path):
    records = [{"Name": "Ada", "City": "Paris"}]
    state = DocumentState(title="Letter")
    first = merge_to_directory(document, state, records, tmp_path)
    second = merge_to_directory(document, state, records, tmp_path)
    assert first[0].exists()
    assert second[0].exists()
    assert first[0] != second[0]


def test_merge_track_changes_disabled(document, tmp_path):
    records = [{"Name": "Ada", "City": "Paris"}]
    state = DocumentState(title="Letter")
    state.track_changes = True
    outputs = merge_to_directory(document, state, records, tmp_path)
    loaded, _ = read_docx(outputs[0])
    assert loaded.track_changes is False


def test_merge_progress_and_cancel(document, tmp_path):
    records = [{"Name": f"P{i}", "City": "C"} for i in range(5)]
    state = DocumentState(title="t")
    seen = []
    outputs = merge_to_directory(
        document, state, records, tmp_path,
        progress=lambda d, t: seen.append(d),
        cancelled=lambda: len(seen) >= 2)
    assert len(outputs) == 2


def test_merge_missing_field_in_source(document, tmp_path):
    records = [{"Name": "Ada"}]
    with pytest.raises(MergeError):
        merge_to_directory(document, DocumentState(title="t"),
                           records, tmp_path)
