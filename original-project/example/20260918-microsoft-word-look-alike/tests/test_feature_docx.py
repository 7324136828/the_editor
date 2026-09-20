import base64
from datetime import datetime
from io import BytesIO
from zipfile import ZipFile

import docx
import pytest
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from lxml import etree
from PySide6.QtGui import QColor, QImage, QTextBlockFormat, QTextCharFormat, QTextCursor, QTextFormat, QTextLength

from folio.design import DesignSettings
from folio.document_features import (
    attachment_node, bookmark, insert_node, insert_section, refresh_fields,
)
from folio.docx_io import DocxError, _import_body, read_docx, write_docx
from folio.editor import SafeDocument
from folio.formatting import paragraph_border
from folio.models import DocumentState, PageSettings


def select(editor, start, end):
    cursor = editor.textCursor()
    cursor.setPosition(start)
    cursor.setPosition(end, QTextCursor.MoveMode.KeepAnchor)
    editor.setTextCursor(cursor)


def export(editor, tmp_path, state=None):
    path = tmp_path / "features.docx"
    write_docx(path, state or DocumentState(), editor.document())
    return path


def xml(path, member="word/document.xml"):
    with ZipFile(path) as archive:
        return etree.fromstring(archive.read(member))


def test_image_description_is_exported_and_imported(editor, tmp_path):
    image = QImage(30, 20, QImage.Format.Format_ARGB32)
    image.fill(QColor("blue"))
    editor.insert_image(image)
    select(editor, 0, 1)
    fmt = QTextCharFormat()
    fmt.setProperty(QTextFormat.Property.ImageAltText, "A blue rectangle")
    editor.textCursor().mergeCharFormat(fmt)
    path = export(editor, tmp_path)
    assert xml(path).find(".//" + qn("wp:docPr")).get("descr") == "A blue rectangle"
    state, _warnings = _import_body(path.read_bytes(), path.name)
    restored = SafeDocument()
    restored.setHtml(state.html)
    fragment = restored.begin().begin().fragment()
    assert fragment.charFormat().property(QTextFormat.Property.ImageAltText) == "A blue rectangle"


def test_scripts_paragraph_adornment_export_and_visible_import(editor, tmp_path):
    editor.setPlainText("H2O and x2")
    select(editor, 1, 2)
    editor.toggle_subscript()
    select(editor, 9, 10)
    editor.toggle_superscript()
    editor.set_paragraph_shading(QColor("#fff0bb"))
    editor.set_paragraph_border("box", "#123456", 2)
    path = export(editor, tmp_path)
    root = xml(path)
    alignments = [element.get(qn("w:val")) for element in root.iter(qn("w:vertAlign"))]
    assert alignments == ["subscript", "superscript"]
    assert root.find(".//" + qn("w:shd")).get(qn("w:fill")) == "#fff0bb"[1:]
    assert len(root.find(".//" + qn("w:pBdr"))) == 4
    state, _warnings = _import_body(path.read_bytes(), path.name)
    restored = SafeDocument()
    restored.setHtml(state.html)
    assert paragraph_border(restored.begin())["color"] == "#123456"
    cursor = QTextCursor(restored)
    cursor.setPosition(1)
    cursor.setPosition(2, QTextCursor.MoveMode.KeepAnchor)
    assert cursor.charFormat().verticalAlignment() == QTextCharFormat.VerticalAlignment.AlignSubScript


@pytest.mark.parametrize("kind, expected", [
    ("next", WD_SECTION_START.NEW_PAGE),
    ("continuous", WD_SECTION_START.CONTINUOUS),
    ("even", WD_SECTION_START.EVEN_PAGE),
    ("odd", WD_SECTION_START.ODD_PAGE),
])
def test_native_section_geometry_and_break_kind(editor, tmp_path, kind, expected):
    state = DocumentState()
    editor.setPlainText("Portrait section")
    cursor = editor.textCursor()
    cursor.movePosition(QTextCursor.MoveOperation.End)
    editor.setTextCursor(cursor)
    page = PageSettings(paper="A4", landscape=True, left_mm=17, right_mm=18, columns=2)
    insert_section(editor, state, kind, page)
    editor.textCursor().insertText("Landscape section")
    path = export(editor, tmp_path, state)
    exported = docx.Document(path)
    assert len(exported.sections) == 2
    assert exported.sections[0].page_width.mm < exported.sections[0].page_height.mm
    second = exported.sections[1]
    assert second.start_type == expected
    assert second.page_width.mm > second.page_height.mm
    assert second.left_margin.mm == pytest.approx(17, abs=0.05)
    assert second._sectPr.find(qn("w:cols")).get(qn("w:num")) == "2"
    assert "\u2063" not in "".join(paragraph.text for paragraph in exported.paragraphs)
    loaded, warnings = read_docx(path)
    assert warnings == []
    assert loaded.features["sections"] == state.features["sections"]


def test_native_bookmark_cross_reference_date_and_signature_controls(editor, tmp_path):
    editor.setPlainText("Heading")
    state = DocumentState()
    editor.selectAll()
    target = bookmark(editor, state, "Introduction")
    cursor = editor.textCursor()
    cursor.movePosition(QTextCursor.MoveOperation.End)
    cursor.insertBlock()
    editor.setTextCursor(cursor)
    insert_node(editor, state, {"kind": "reference", "target": target, "display": "text"}, "Heading")
    editor.textCursor().insertText(" ")
    insert_node(editor, state, {"kind": "datetime", "format": "date"}, "date")
    editor.textCursor().insertText(" ")
    insert_node(editor, state, {"kind": "signature", "name": "Recipient"}, "Signature: Recipient")
    refresh_fields(editor.document(), state.features, datetime(2026, 9, 19, 13, 15))
    path = export(editor, tmp_path, state)
    root = xml(path)
    bookmarks = {node.get(qn("w:name")) for node in root.iter(qn("w:bookmarkStart"))}
    fields = [node.get(qn("w:instr")) for node in root.iter(qn("w:fldSimple"))]
    reference = next(field for field in fields if field.startswith("REF "))
    assert reference.split()[1] in bookmarks
    assert any(field.startswith('DATE \\@ "MMMM dd, yyyy"') for field in fields)
    assert "September 19, 2026" in "".join(root.itertext())
    control = root.find(".//" + qn("w:sdt"))
    assert control is not None
    assert "Signature: Recipient" in "".join(control.itertext())
    imported, _warnings = _import_body(path.read_bytes(), path.name)
    restored = SafeDocument()
    restored.setHtml(imported.html)
    assert "September 19, 2026" in restored.toPlainText()
    assert "Signature: Recipient" in restored.toPlainText()


def test_attachments_are_byte_exact_packages(editor, tmp_path):
    attachment = tmp_path / "sample.bin"
    payload = b"Folio attachment\x00\x01\xff"
    attachment.write_bytes(payload)
    state = DocumentState()
    insert_node(editor, state, attachment_node(attachment), "Attachment: sample.bin")
    path = export(editor, tmp_path, state)
    with ZipFile(path) as archive:
        names = [name for name in archive.namelist() if name.startswith("word/embeddings/")]
        assert len(names) == 1
        assert archive.read(names[0]) == payload
        relationships = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
        package = next(node for node in relationships if node.get("Type", "").endswith("/package"))
        assert "word/" + package.get("Target") == names[0]
    loaded, warnings = read_docx(path)
    assert warnings == []
    node = next(iter(loaded.features["nodes"].values()))
    assert base64.b64decode(node["data"]) == payload


def test_floating_text_is_anchored_and_hidden_objects_omitted(editor, tmp_path):
    editor.setPlainText("Body text")
    state = DocumentState()
    state.features["objects"] = [
        {"id": "one", "kind": "text", "name": "Note", "text": "Floating text",
         "x": 90, "y": 100, "width": 180, "height": 90, "wrap": "square"},
        {"id": "two", "kind": "text", "text": "Hidden text", "visible": False},
    ]
    path = export(editor, tmp_path, state)
    root = xml(path)
    anchors = list(root.iter(qn("wp:anchor")))
    assert len(anchors) == 1
    anchor = anchors[0]
    assert anchor.get("behindDoc") == "0"
    assert anchor.find(qn("wp:wrapSquare")) is not None
    horizontal = anchor.find(qn("wp:positionH"))
    assert horizontal.get("relativeFrom") == "page"
    assert int(horizontal.find(qn("wp:posOffset")).text) == 90 * 9525


def test_theme_page_decoration_and_header_aliases_export(editor, tmp_path):
    editor.setPlainText("Themed title")
    settings = DesignSettings(theme="Woodland", style_set="Compact", page_color="#fcf5dd",
                              page_border="double", watermark_kind="text", watermark_text="DRAFT")
    editor.design_settings = settings
    editor.apply_style("Title")
    state = DocumentState(design=settings.to_dict())
    state.page.header = "Page {CurrentPage} of {TotalPages}"
    path = export(editor, tmp_path, state)
    exported = docx.Document(path)
    assert exported.styles["Title"].font.name == "Cambria"
    assert exported.styles["Title"].font.size.pt == 26
    assert exported.paragraphs[0].style.name == "Title"
    root = xml(path)
    assert root.find(qn("w:background")).get(qn("w:color")) == "fcf5dd"
    assert root.find(".//" + qn("w:pgBorders") + "/" + qn("w:top")).get(qn("w:val")) == "double"
    header = exported.sections[0].header._element
    fields = [node.text.strip() for node in header.iter(qn("w:instrText"))]
    assert fields == ["PAGE", "NUMPAGES"]
    assert len(list(header.iter(qn("wp:anchor")))) == 1
    theme = xml(path, "word/theme/theme1.xml")
    major = theme.find(".//" + qn("a:majorFont") + "/" + qn("a:latin"))
    assert major.get("typeface") == "Cambria"


def test_table_dimensions_are_native_word_widths_and_minimum_heights(editor, tmp_path):
    table = editor.insert_table(2, 2)
    fmt = table.format()
    fmt.setColumnWidthConstraints([QTextLength(QTextLength.Type.FixedLength, width)
                                   for width in (120, 180)])
    table.setFormat(fmt)
    cursor = table.cellAt(0, 0).firstCursorPosition()
    block = cursor.blockFormat()
    block.setLineHeight(40, QTextBlockFormat.LineHeightTypes.MinimumHeight.value)
    cursor.setBlockFormat(block)
    path = export(editor, tmp_path)
    exported = docx.Document(path).tables[0]
    assert exported.columns[0].width == 120 * 9525
    assert exported.columns[1].width == 180 * 9525
    assert exported.rows[0].height == 40 * 9525
    assert exported.rows[0]._tr.trPr.find(qn("w:trHeight")).get(qn("w:hRule")) == "atLeast"


@pytest.mark.parametrize("change", ["design", "features"])
def test_invalid_new_metadata_does_not_replace_existing_file(editor, tmp_path, change):
    path = tmp_path / "existing.docx"
    path.write_bytes(b"existing")
    state = DocumentState()
    if change == "design":
        state.design = {"page_color": "invalid"}
    else:
        state.features = {"objects": [{"id": "x", "kind": "image", "width": -10}]}
    with pytest.raises(DocxError):
        write_docx(path, state, editor.document())
    assert path.read_bytes() == b"existing"
