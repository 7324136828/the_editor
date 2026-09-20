import zipfile
from xml.etree import ElementTree

import docx
import pytest
from PySide6.QtGui import QImage, QPainter, QTextCharFormat, QTextCursor

from folio.document_features import (checked_features, clipboard_feature_payload, insert_node,
                                     node_ranges, prepare_pasted_nodes, refresh_fields)
from folio.editor import SafeDocument
from folio.docx_io import read_docx, write_docx
from folio.models import DocumentState, PageSettings, ValidationError
from folio.publishing import Publication
from folio.references import (add_source, bibliography_text, citation_text, ensure_collection,
                              insert_citation, insert_note, note_entries, reference_values,
                              refresh_references, registry, tag_selection)


def select(editor, start, end=None):
    cursor = editor.textCursor()
    cursor.setPosition(start)
    cursor.setPosition(start if end is None else end, QTextCursor.MoveMode.KeepAnchor)
    editor.setTextCursor(cursor)
    return cursor


def append(editor, text="\n"):
    cursor = editor.textCursor()
    cursor.movePosition(QTextCursor.MoveOperation.End)
    cursor.setCharFormat(QTextCharFormat())
    cursor.insertText(text)
    editor.setTextCursor(cursor)


def value(document, key):
    start, end = node_ranges(document)[key]
    cursor = QTextCursor(document)
    cursor.setPosition(start)
    cursor.setPosition(end, QTextCursor.MoveMode.KeepAnchor)
    return cursor.selectedText().replace("\u2029", "\n")


def test_citation_styles_and_unicode_source_roundtrip(editor):
    state = DocumentState()
    source = {"title": "Research 🧬", "author": "Nguyen, An; Smith, Jo", "year": "2024", "publisher": "Lab Press"}
    key = add_source(state, source)
    citation = insert_citation(editor, state, key, "23–25")
    append(editor)
    bibliography = ensure_collection(editor, state, "bibliography")
    refresh_references(editor.document(), state)
    assert value(editor.document(), citation) == "(Nguyen & Smith, 2024, p. 23–25)"
    assert "Research 🧬" in value(editor.document(), bibliography)
    registry(state)["style"] = "IEEE"
    refresh_references(editor.document(), state)
    assert value(editor.document(), citation) == "[1, p. 23–25]"
    assert "[1]" in value(editor.document(), bibliography)
    state.html = editor.document().toHtml()
    restored = DocumentState.from_dict(state.to_dict())
    document = SafeDocument()
    document.setHtml(restored.html)
    assert node_ranges(document)[citation] == node_ranges(editor.document())[citation]
    assert restored.features["references"]["sources"][key]["title"] == "Research 🧬"


def test_citation_numbering_follows_first_use_and_deduplicates(editor):
    state = DocumentState()
    registry(state)["style"] = "IEEE"
    a = add_source(state, {"title": "A"})
    b = add_source(state, {"title": "B"})
    first = insert_citation(editor, state, b)
    append(editor, " ")
    second = insert_citation(editor, state, a)
    append(editor, " ")
    third = insert_citation(editor, state, b)
    append(editor)
    bibliography = ensure_collection(editor, state, "bibliography")
    refresh_references(editor.document(), state)
    assert [value(editor.document(), key) for key in (first, second, third)] == ["[1]", "[2]", "[1]"]
    assert value(editor.document(), bibliography).count('"B,"') == 1
    start, end = node_ranges(editor.document())[first]
    select(editor, start, end).removeSelectedText()
    refresh_references(editor.document(), state)
    assert [value(editor.document(), key) for key in (second, third)] == ["[1]", "[2]"]


def test_supported_common_citation_formats():
    source = {"author": "Rivera, Ada", "title": "Example", "year": "2025"}
    assert citation_text(source, "MLA", locator="8") == "(Rivera 8)"
    assert citation_text(source, "Chicago", locator="8") == "(Rivera 2025, 8)"
    assert bibliography_text(source, "APA") == "Rivera, Ada (2025). Example."


def test_notes_have_separate_sequences_and_endnote_container(editor):
    state = DocumentState()
    editor.setPlainText("Text 🧬")
    select(editor, 7)
    first = insert_note(editor, state, "First footnote")
    append(editor, " text ")
    terminal = insert_note(editor, state, "Terminal note", "endnote")
    assert editor.textCursor().position() == node_ranges(editor.document())[terminal][1]
    select(editor, 0)
    prior = insert_note(editor, state, "Inserted before")
    refresh_references(editor.document(), state)
    assert value(editor.document(), prior) == "1"
    assert value(editor.document(), first) == "2"
    assert value(editor.document(), terminal) == "1"
    assert "Endnotes\n1. Terminal note" in editor.toPlainText()
    assert [entry[2] for entry in note_entries(editor.document(), state, "footnote")] == [1, 2]


def test_captions_lists_and_page_references_update_after_deletion(editor):
    state = DocumentState()
    first = insert_node(editor, state, {"kind": "caption", "label": "Figure", "text": "Alpha 🧬"}, "Caption")
    append(editor)
    second = insert_node(editor, state, {"kind": "caption", "label": "Figure", "text": "Beta"}, "Caption")
    append(editor)
    table = insert_node(editor, state, {"kind": "caption", "label": "Table", "text": "Data"}, "Caption")
    append(editor)
    listing = ensure_collection(editor, state, "figures", "Figure")
    append(editor)
    crossref = insert_node(editor, state, {"kind": "reference", "target": second, "display": "text"}, "Reference")
    refresh_references(editor.document(), state, lambda position: 7)
    refresh_fields(editor.document(), state.features, page_for_position=lambda position: 7)
    assert value(editor.document(), second) == "Figure 2: Beta"
    assert value(editor.document(), table) == "Table 1: Data"
    assert "p. 7" in value(editor.document(), listing)
    assert "Data" not in value(editor.document(), listing)
    start, end = node_ranges(editor.document())[first]
    select(editor, start, end).removeSelectedText()
    refresh_references(editor.document(), state, lambda position: 3)
    refresh_fields(editor.document(), state.features)
    assert value(editor.document(), crossref) == "Figure 1: Beta"


def test_index_and_authorities_use_anchored_utf16_ranges(editor):
    state = DocumentState()
    editor.setPlainText("🧬 apple banana apple Case X")
    select(editor, 3, 8)
    first = tag_selection(editor, state, "apple")
    select(editor, 9, 15)
    tag_selection(editor, state, "Banana")
    select(editor, 16, 21)
    tag_selection(editor, state, "apple")
    select(editor, 22, 28)
    tag_selection(editor, state, "Case X", "Cases")
    append(editor)
    index = ensure_collection(editor, state, "index")
    authorities = ensure_collection(editor, state, "authorities")
    refresh_references(editor.document(), state, lambda position: 2 if position > 15 else 1)
    assert "apple  1, 2\nBanana  1" in value(editor.document(), index)
    assert "Cases\nCase X  2" in value(editor.document(), authorities)
    select(editor, 0).insertText("🧬")
    assert node_ranges(editor.document())[first] == (5, 10)


@pytest.mark.parametrize("features", [
    {"references": {"style": "unsupported"}},
    {"references": {"sources": {"source": {"title": ""}}}},
    {"nodes": {"note": {"kind": "note", "type": "footnote", "text": 123}}},
    {"nodes": {"caption": {"kind": "caption", "label": "Unknown", "text": "A"}}},
    {"nodes": {"list": {"kind": "generated", "collection": "code"}}},
    {"nodes": {"term": {"kind": "index_entry", "term": ""}}},
])
def test_invalid_reference_metadata_is_rejected(features):
    with pytest.raises(ValidationError):
        checked_features(features)


def test_copied_citations_keep_source_data(editor):
    state = DocumentState()
    source = add_source(state, {"title": "Portable", "author": "Nguyen, An", "year": "2026"})
    insert_citation(editor, state, source)
    document = SafeDocument()
    document.setHtml(editor.document().toHtml())
    target = DocumentState()
    payload = clipboard_feature_payload(document, state)
    prepare_pasted_nodes(document, target, payload)
    assert len(target.features["nodes"]) == 1
    refresh_references(document, target)
    assert document.toPlainText() == "(Nguyen, 2026)"
    values = reference_values(document, target, lambda position: 1)
    assert "[Source missing]" not in values.values()


def test_footnotes_reserve_body_space_and_render_without_modifying_source(editor):
    state = DocumentState()
    editor.setPlainText("A short body.")
    select(editor, editor.document().characterCount() - 1)
    note = insert_note(editor, state, "A published footnote.")
    editor.document().folio_state = state
    before = editor.document().toHtml()
    publication = Publication(editor.document(), state.page, "Notes")
    assert publication._footnote_height > 0
    assert 0 in publication._footnotes
    assert publication.page_for_position(node_ranges(editor.document())[note][0]) == 1
    image = QImage(int(publication.page_w), int(publication.page_h), QImage.Format.Format_ARGB32)
    image.fill(0xffffffff)
    painter = QPainter(image)
    try:
        publication.paint_page(painter, 0)
    finally:
        painter.end()
    assert editor.document().toHtml() == before
    assert any(image.pixelColor(x, y).value() < 200
               for y in range(int(publication.margin_top + publication.body_h + 18),
                              int(publication.page_h - publication.margin_bottom))
               for x in range(int(publication.margin_left), int(publication.page_w - publication.margin_right), 3))


def test_long_footnotes_continue_on_following_pages(editor):
    state = DocumentState()
    insert_note(editor, state, "Long note. " * 700)
    publication = Publication(editor.document(), state.page, "Long notes", state)
    assert len(publication._footnotes) > 1
    assert publication.page_count() > publication.body_page_count()
    for entries in publication._footnotes.values():
        assert all(offset + height <= publication._footnote_height - 18 + 0.1
                   for document, source_y, height, offset in entries)


def test_native_footnotes_and_metadata_roundtrip(editor, tmp_path):
    state = DocumentState()
    editor.setPlainText("A body 🧬")
    select(editor, editor.document().characterCount() - 1)
    key = insert_note(editor, state, "A native footnote.\nSecond line.")
    path = tmp_path / "notes.docx"
    write_docx(path, state, editor.document())
    with zipfile.ZipFile(path) as archive:
        notes = ElementTree.fromstring(archive.read("word/footnotes.xml"))
        body = ElementTree.fromstring(archive.read("word/document.xml"))
        relationships = archive.read("word/_rels/document.xml.rels").decode()
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        assert len(body.findall(".//w:footnoteReference", ns)) == 1
        assert "footnotes.xml" in relationships
        assert "A native footnote." in "".join(notes.itertext())
        assert "Second line." in "".join(notes.itertext())
    restored, warnings = read_docx(path)
    assert not warnings
    assert restored.features["nodes"][key]["text"] == "A native footnote.\nSecond line."
    document = SafeDocument()
    document.setHtml(restored.html)
    assert key in node_ranges(document)


def test_external_word_edit_preserves_native_footnote_text(editor, tmp_path):
    state = DocumentState()
    editor.setPlainText("Before ")
    select(editor, editor.document().characterCount() - 1)
    insert_note(editor, state, "Native note 🧬 text.\nSecond paragraph.")
    path = tmp_path / "external-note.docx"
    write_docx(path, state, editor.document())
    external = docx.Document(path)
    external.add_paragraph("External change")
    external.save(path)
    restored, warnings = read_docx(path)
    assert warnings
    notes = [node for node in restored.features["nodes"].values() if node["kind"] == "note"]
    assert len(notes) == 1
    assert notes[0]["text"] == "Native note 🧬 text.\nSecond paragraph."
    document = SafeDocument()
    document.setHtml(restored.html)
    assert "External change" in document.toPlainText()
    assert len(note_entries(document, restored)) == 1


def test_reference_ribbon_notes_navigation_and_readonly(qapp, store, monkeypatch):
    from PySide6.QtWidgets import QInputDialog
    from folio.window import FolioWindow

    window = FolioWindow(store=store, restore=False)
    try:
        window.load_state(DocumentState(html="<p>Body</p>"), clean=True)
        monkeypatch.setattr(QInputDialog, "getMultiLineText", lambda *args, **kwargs: ("Editable note", True))
        window._insert_reference_note("footnote")
        assert window.notes_list.count() == 1
        assert window.note_text.toPlainText() == "Editable note"
        window.note_text.setPlainText("Changed note")
        window._save_note()
        window._goto_note_anchor()
        key = next(key for key, node in window.state.features["nodes"].items() if node["kind"] == "note")
        assert window.editor.textCursor().selectionStart() == node_ranges(window.editor.document())[key][0]
        assert window.state.features["nodes"][key]["text"] == "Changed note"
        window._permission_mode("Viewing")
        window._insert_reference_note("footnote")
        assert window.notes_list.count() == 1
    finally:
        window._dirty = False
        window.close()
        qapp.processEvents()


def test_source_manager_copies_document_sources_to_local_library(qapp, store, monkeypatch):
    from PySide6.QtWidgets import QDialog, QListWidget, QPushButton
    from folio.window import FolioWindow

    window = FolioWindow(store=store, restore=False)
    try:
        key = add_source(window.state, {"title": "Reusable source", "author": "Lee, Ada"})

        def manage(dialog):
            listing = dialog.findChild(QListWidget)
            listing.setCurrentRow(0)
            button = next(button for button in dialog.findChildren(QPushButton) if button.text() == "Copy to Other Library")
            button.click()
            return QDialog.DialogCode.Rejected

        monkeypatch.setattr(QDialog, "exec", manage)
        window._manage_sources()
        assert store.get_setting("reference_sources")[key]["title"] == "Reusable source"
    finally:
        window._dirty = False
        window.close()
        qapp.processEvents()


def test_automatic_reference_updates_respect_protected_ranges(qapp, store):
    from folio.window import FolioWindow

    window = FolioWindow(store=store, restore=False)
    try:
        window.editor.clear()
        key = insert_node(window.editor, window.state,
                          {"kind": "caption", "label": "Figure", "text": "Original"}, "Figure 1: Original")
        append(window.editor, " trailing")
        start, end = node_ranges(window.editor.document())[key]
        window.state.review["locks"] = [{"id": "range", "start": start, "end": end, "name": "Protected caption"}]
        window.tracker.refresh_protection()
        window.state.features["nodes"][key]["text"] = "New metadata"
        before = window.editor.document().toHtml()
        window._refresh_references()
        assert window.editor.document().toHtml() == before
        assert value(window.editor.document(), key) == "Figure 1: Original"
        window.state.review["locks"] = []
        window.state.review["formatting_locked"] = True
        window.tracker.refresh_protection()
        window._refresh_references()
        assert window.editor.document().toHtml() == before
        window.state.review["formatting_locked"] = False
        window.tracker.refresh_protection()
        window._refresh_references()
        assert value(window.editor.document(), key) == "Figure 1: New metadata"
    finally:
        window._dirty = False
        window.close()
        qapp.processEvents()


def test_footnotes_keep_document_numbering_across_sections(editor):
    from folio.document_features import insert_section

    state = DocumentState()
    insert_note(editor, state, "First section note")
    insert_section(editor, state, "next", PageSettings(columns=2))
    append(editor, "Second section body ")
    insert_note(editor, state, "Second section note")
    refresh_references(editor.document(), state)
    publication = Publication(editor.document(), state.page, "Sections", state)
    assert publication.page_count() == 2
    assert len(publication._sections) == 2
    child = publication._sections[1][4]
    assert child._footnotes[0][0][0].toPlainText().startswith("2. Second section note")


@pytest.mark.parametrize("kind", ["note", "caption", "citation"])
@pytest.mark.parametrize("typing", ["keys", "plain"])
@pytest.mark.parametrize("reposition", ["none", "end", "start"])
def test_prose_after_reference_field_survives_refresh(editor, kind, typing, reposition):
    from PySide6.QtTest import QTest

    state = DocumentState()
    if kind == "note":
        key = insert_note(editor, state, "Note text")
    elif kind == "citation":
        source = add_source(state, {"author": "Lee, Ada", "title": "Example", "year": "2026"})
        key = insert_citation(editor, state, source)
    else:
        key = insert_node(editor, state, {"kind": "caption", "label": "Figure", "text": "Example"}, "Figure 1: Example")
    original = value(editor.document(), key)
    if reposition != "none":
        select(editor, 0)
        select(editor, node_ranges(editor.document())[key][1 if reposition == "end" else 0])
    if typing == "keys":
        QTest.keyClicks(editor, " prose after the field")
    else:
        editor.insertPlainText(" prose after the field")
    refresh_references(editor.document(), state)
    assert " prose after the field" in editor.toPlainText()
    assert value(editor.document(), key) == original
