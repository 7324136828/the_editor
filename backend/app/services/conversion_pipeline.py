"""Document conversion and extraction pipeline with isolated system temp execution.

Supports Word (.docx), Excel (.xlsx), PowerPoint (.pptx), PDF (.pdf), CSV, and text files.
Generates extracted text, markdown representation, and metadata, packaged into a ZIP archive.
"""

from __future__ import annotations

import json
import traceback
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from ..schemas.job import JobRecord
from ..utils.temp_manager import TempJobManager


class ConversionPipeline:
    """Handles staging, extraction, conversion, and archive creation."""

    def __init__(self, temp_manager: TempJobManager) -> None:
        self.temp_manager = temp_manager

    def process_job(self, job: JobRecord, file_path: Path, job_dir: Path) -> None:
        """Execute document conversion pipeline and produce ZIP package."""
        outputs_dir = job_dir / "outputs"
        ext = file_path.suffix.lower()

        extracted_text = ""
        markdown_content = ""
        metadata: dict[str, Any] = {
            "original_filename": job.filename,
            "file_size": job.file_size,
            "processed_at": datetime.now(timezone.utc).isoformat(),
            "format": ext.lstrip("."),
        }

        try:
            job.status = "in_progress"
            job.progress = 25

            if ext == ".docx":
                extracted_text, markdown_content, meta = self._process_docx(file_path)
                metadata.update(meta)
            elif ext in (".xlsx", ".xls"):
                extracted_text, markdown_content, meta = self._process_xlsx(file_path)
                metadata.update(meta)
            elif ext == ".pptx":
                extracted_text, markdown_content, meta = self._process_pptx(file_path)
                metadata.update(meta)
            elif ext == ".pdf":
                extracted_text, markdown_content, meta = self._process_pdf(file_path)
                metadata.update(meta)
            elif ext in (".csv", ".txt", ".md", ".json", ".ts", ".js", ".py"):
                raw = file_path.read_text(encoding="utf-8", errors="replace")
                extracted_text = raw
                markdown_content = f"# {job.filename}\n\n```\n{raw}\n```"
                metadata["characters"] = len(raw)
                metadata["lines"] = len(raw.splitlines())
            else:
                raw_bytes = file_path.read_bytes()
                extracted_text = f"Binary file: {job.filename} ({len(raw_bytes)} bytes)"
                markdown_content = f"# {job.filename}\n\nBinary file preview not available."

            job.progress = 70

            # Write outputs
            (outputs_dir / "extracted_text.txt").write_text(extracted_text, encoding="utf-8")
            (outputs_dir / "content.md").write_text(markdown_content, encoding="utf-8")
            (outputs_dir / "metadata.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")

            # Package ZIP
            job.progress = 90
            zip_path = self.temp_manager.package_zip(job_dir, f"{Path(job.filename).stem}_converted.zip")

            job.status = "completed"
            job.progress = 100
            job.completed_at = datetime.now(timezone.utc).isoformat()
            job.zip_path = str(zip_path)
            job.summary = {
                "format": ext.lstrip("."),
                "text_length": len(extracted_text),
                "metadata": metadata,
            }

        except Exception as e:
            job.status = "failed"
            job.error_message = f"Conversion failed: {e}\n{traceback.format_exc()}"
            job.completed_at = datetime.now(timezone.utc).isoformat()

    def _process_docx(self, path: Path) -> tuple[str, str, dict[str, Any]]:
        text_lines = []
        md_lines = [f"# Document: {path.stem}\n"]
        try:
            import docx
            doc = docx.Document(path)
            for p in doc.paragraphs:
                if not p.text.strip():
                    continue
                text_lines.append(p.text)
                if p.style.name.startswith("Heading 1"):
                    md_lines.append(f"# {p.text}\n")
                elif p.style.name.startswith("Heading 2"):
                    md_lines.append(f"## {p.text}\n")
                elif p.style.name.startswith("Heading 3"):
                    md_lines.append(f"### {p.text}\n")
                else:
                    md_lines.append(f"{p.text}\n")
            words = sum(len(line.split()) for line in text_lines)
            return "\n".join(text_lines), "\n".join(md_lines), {"words": words, "paragraphs": len(doc.paragraphs)}
        except Exception:
            # Fallback if docx module fails or invalid format
            raw = path.read_text(encoding="utf-8", errors="ignore")
            return raw, f"```\n{raw}\n```", {}

    def _process_xlsx(self, path: Path) -> tuple[str, str, dict[str, Any]]:
        text_lines = []
        md_lines = [f"# Workbook: {path.stem}\n"]
        sheet_count = 0
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True)
            sheet_count = len(wb.sheetnames)
            for name in wb.sheetnames:
                sheet = wb[name]
                text_lines.append(f"--- Sheet: {name} ---")
                md_lines.append(f"## Sheet: {name}\n")
                rows = list(sheet.iter_rows(values_only=True))
                if rows:
                    header = rows[0]
                    md_lines.append("| " + " | ".join(str(c or "") for c in header) + " |")
                    md_lines.append("| " + " | ".join(["---"] * len(header)) + " |")
                    for row in rows[1:100]:
                        row_vals = [str(c) if c is not None else "" for c in row]
                        text_lines.append("\t".join(row_vals))
                        md_lines.append("| " + " | ".join(row_vals) + " |")
                    md_lines.append("\n")
            return "\n".join(text_lines), "\n".join(md_lines), {"sheets": wb.sheetnames, "sheet_count": sheet_count}
        except Exception:
            raw = path.read_text(encoding="utf-8", errors="ignore")
            return raw, f"```\n{raw}\n```", {}

    def _process_pptx(self, path: Path) -> tuple[str, str, dict[str, Any]]:
        text_lines = []
        md_lines = [f"# Presentation: {path.stem}\n"]
        slide_count = 0
        try:
            import pptx
            prs = pptx.Presentation(path)
            slide_count = len(prs.slides)
            for i, slide in enumerate(prs.slides, 1):
                title = slide.shapes.title.text if slide.shapes.title else f"Slide {i}"
                text_lines.append(f"--- Slide {i}: {title} ---")
                md_lines.append(f"## Slide {i}: {title}\n")
                for shape in slide.shapes:
                    if shape.has_text_frame and shape != slide.shapes.title:
                        for p in shape.text_frame.paragraphs:
                            if p.text.strip():
                                text_lines.append(f"  • {p.text}")
                                md_lines.append(f"- {p.text}")
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        text_lines.append(f"  Notes: {notes}")
                        md_lines.append(f"\n> **Notes**: {notes}\n")
                md_lines.append("\n")
            return "\n".join(text_lines), "\n".join(md_lines), {"slides_count": slide_count}
        except Exception:
            raw = path.read_text(encoding="utf-8", errors="ignore")
            return raw, f"```\n{raw}\n```", {}

    def _process_pdf(self, path: Path) -> tuple[str, str, dict[str, Any]]:
        text_lines = []
        md_lines = [f"# Document: {path.stem}\n"]
        page_count = 0
        try:
            import pypdf
            reader = pypdf.PdfReader(path)
            page_count = len(reader.pages)
            for i, page in enumerate(reader.pages, 1):
                t = page.extract_text() or ""
                text_lines.append(f"--- Page {i} ---\n{t}")
                md_lines.append(f"## Page {i}\n\n{t}\n")
            words = sum(len(line.split()) for line in text_lines)
            return "\n".join(text_lines), "\n".join(md_lines), {"pages": page_count, "words": words}
        except Exception:
            raw = path.read_text(encoding="utf-8", errors="ignore")
            return raw, f"```\n{raw}\n```", {}

